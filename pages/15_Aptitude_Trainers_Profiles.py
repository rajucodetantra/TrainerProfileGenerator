import os
import re
import sys
import shutil
import subprocess
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import streamlit as st

st.set_page_config(
    page_title="Aptitude Trainer Profiles",
    page_icon="🧠",
    layout="wide",
    initial_sidebar_state="expanded",
)

ROOT_DIR = Path(__file__).resolve().parent.parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

# Existing login/styles are only used; they are not modified.
try:
    from login import require_login
    require_login()
except ImportError:
    st.warning("login.py not found. Page opened without login protection.")

try:
    from ui_styles import apply_global_styles
    apply_global_styles()
except Exception:
    pass

from openpyxl import load_workbook
from PIL import Image as PILImage
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# =========================================================
# FIXED FILES USED ONLY BY THIS PAGE
# =========================================================

EXCEL_PATH = ROOT_DIR / "data" / "Aptitude_Trainer.xlsx"
TEMPLATE_PATH = ROOT_DIR / "templates" / "Aptitude_Trainer_Profile_Template.docx"
OUTPUT_DIR = ROOT_DIR / "output" / "Aptitude_Trainer_Profiles"
WORD_OUTPUT_DIR = OUTPUT_DIR / "Word"
PDF_OUTPUT_DIR = OUTPUT_DIR / "PDF"
TEMP_IMAGE_DIR = OUTPUT_DIR / "_temp_images"

for folder in [OUTPUT_DIR, WORD_OUTPUT_DIR, PDF_OUTPUT_DIR, TEMP_IMAGE_DIR]:
    folder.mkdir(parents=True, exist_ok=True)

EXPECTED_FIELDS = [
    "S.No",
    "Name",
    "Employee ID",
    "Designation",
    "Company Mail",
    "Phone Number",
    "Date of Joining",
    "Today Date",
    "Experience",
    "Qualification",
    "Certifications",
    "Skills",
    "Photo",
]

COLUMN_ALIASES = {
    "sno": "S.No",
    "s.no": "S.No",
    "s.no.": "S.No",
    "serialno": "S.No",
    "serialnumber": "S.No",
    "name": "Name",
    "trainername": "Name",
    "employeeid": "Employee ID",
    "empid": "Employee ID",
    "designation": "Designation",
    "companymail": "Company Mail",
    "companyemail": "Company Mail",
    "phonenumber": "Phone Number",
    "phone": "Phone Number",
    "mobile": "Phone Number",
    "mobilenumber": "Phone Number",
    "dateofjoining": "Date of Joining",
    "doj": "Date of Joining",
    "todaydate": "Today Date",
    "currentdate": "Today Date",
    "experience": "Experience",
    "experienceinyears": "Experience",
    "qualification": "Qualification",
    "education": "Qualification",
    "certification": "Certifications",
    "certifications": "Certifications",
    "skill": "Skills",
    "skills": "Skills",
    "coreskills": "Skills",
    "photo": "Photo",
    "image": "Photo",
    "photos": "Photo",
}


def clean_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_header(value):
    text = clean_text(value).lower().replace("_", " ")
    text = re.sub(r"\s+", " ", text).strip()
    compact = re.sub(r"[^a-z0-9.]", "", text)
    if text in COLUMN_ALIASES:
        return COLUMN_ALIASES[text]
    if compact in COLUMN_ALIASES:
        return COLUMN_ALIASES[compact]
    return clean_text(value)


def safe_filename(value):
    value = clean_text(value)
    value = re.sub(r'[<>:"/\\|?*]', "", value)
    value = re.sub(r"\s+", "_", value)
    return value or "Trainer"


def format_date_value(value):
    if value is None or value == "":
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%d/%m/%Y")
    return clean_text(value)


def parse_date_value(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean_text(value)
    for fmt in ["%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y", "%d-%m-%y", "%m/%d/%Y"]:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def calculate_experience(date_of_joining, today_value=None):
    start = parse_date_value(date_of_joining)
    if start is None:
        return ""
    end = parse_date_value(today_value) or date.today()
    if end < start:
        return ""
    years = end.year - start.year
    months = end.month - start.month
    if end.day < start.day:
        months -= 1
    if months < 0:
        years -= 1
        months += 12
    parts = []
    if years > 0:
        parts.append(f"{years} Year" + ("s" if years != 1 else ""))
    if months > 0:
        parts.append(f"{months} Month" + ("s" if months != 1 else ""))
    return " ".join(parts) if parts else "Less than 1 Month"


def split_skills(value):
    text = clean_text(value)
    if not text:
        return []
    text = re.sub(r"(?i)<br\s*/?>", "\n", text)
    for symbol in ["■", "●", "•"]:
        text = text.replace(symbol, "\n")
    parts = re.split(r"[\n,;|]+", text)
    result = []
    seen = set()
    for item in parts:
        item = re.sub(r"\s+", " ", item.strip(" \t\r\n-–—*▪")).strip()
        if not item:
            continue
        key = item.lower()
        if key not in seen:
            seen.add(key)
            result.append(item)
    return result


def format_core_skills(skills):
    return "■ " + " ■ ".join(skills) if skills else ""


def bullet_lines(items):
    result = []
    for item in items:
        item = clean_text(item).lstrip("■●•▪-* ").strip()
        if item:
            result.append("● " + item)
    return "\n".join(result)


def build_certifications(value):
    text = clean_text(value)
    if not text:
        return ""
    if text.lower() in ["no", "none", "na", "n/a", "nil"]:
        return "No certifications mentioned."
    text = re.sub(r"(?i)<br\s*/?>", "\n", text)
    return bullet_lines(re.split(r"[\n;|]+", text))


def build_professional_summary(row, skills):
    points = ["Aptitude Trainer focused on structured aptitude and reasoning training."]
    if skills:
        points.append("Core strengths include " + ", ".join(skills[:6]) + ".")
    qualification = clean_text(row.get("Qualification"))
    if qualification:
        points.append(f"Academic qualification: {qualification}.")
    points.append(
        "Emphasizes concept clarity, analytical thinking, problem-solving and student engagement."
    )
    return bullet_lines(points)


def build_professional_highlights(skills):
    lower_skills = [skill.lower() for skill in skills]
    points = []
    if any("aptitude" in item for item in lower_skills):
        points.append("Supports structured aptitude learning with emphasis on concept clarity and problem-solving.")
    if any("reasoning" in item for item in lower_skills):
        points.append("Focuses on logical reasoning and analytical approaches for solving problems efficiently.")
    if any("analytical" in item for item in lower_skills):
        points.append("Encourages analytical thinking through step-by-step problem interpretation.")
    if any("student engagement" in item for item in lower_skills):
        points.append("Uses student-engagement practices to support active participation during training.")
    if any("conceptual" in item for item in lower_skills):
        points.append("Uses conceptual teaching to simplify aptitude and reasoning topics.")
    if not points:
        points = [
            "Focuses on concept clarity and structured problem-solving.",
            "Supports learner participation through clear and practical training.",
        ]
    return bullet_lines(points)


def build_subjects_handled(skills):
    keywords = [
        "aptitude", "quantitative", "logical reasoning", "reasoning",
        "verbal", "data interpretation", "analytical thinking",
        "problem-solving", "problem solving",
    ]
    subjects = [skill for skill in skills if any(k in skill.lower() for k in keywords)]
    if not subjects:
        subjects = skills[:4]
    return bullet_lines(subjects)


def build_training_expertise(skills):
    results = []
    for skill in skills:
        lower = skill.lower()
        if any(k in lower for k in ["aptitude", "reasoning", "analytical", "problem"]):
            results.append(skill if "training" in lower else skill + " Training")
        elif "conceptual" in lower or "student engagement" in lower:
            results.append(skill)
    if not results:
        results = skills[:6]
    return bullet_lines(results)


def build_core_competencies(skills):
    preferred = ["analytical", "problem", "conceptual", "student engagement", "communication", "reasoning", "teaching", "training"]
    results = [skill for skill in skills if any(k in skill.lower() for k in preferred)]
    if not results:
        results = skills[:6]
    return bullet_lines(results)


def build_exam_expertise(skills):
    keywords = ["placement", "crt", "cat", "gre", "ssc", "bank", "competitive exam", "campus recruitment", "government exam"]
    return bullet_lines([skill for skill in skills if any(k in skill.lower() for k in keywords)])


# =========================================================
# EXCEL READING
# =========================================================


def find_data_sheet(workbook):
    for worksheet in workbook.worksheets:
        for row_number in range(1, min(worksheet.max_row, 10) + 1):
            headers = [normalize_header(cell.value) for cell in worksheet[row_number]]
            if "Name" in headers and "Employee ID" in headers:
                return worksheet, row_number
    raise ValueError("Could not find a sheet containing Name and Employee ID headers.")


def prepare_row(source_row):
    row = {field: source_row.get(field, "") for field in EXPECTED_FIELDS}
    row["_excel_row"] = source_row.get("_excel_row")
    row["Name"] = clean_text(row.get("Name"))
    row["Employee ID"] = clean_text(row.get("Employee ID"))
    row["Designation"] = clean_text(row.get("Designation")) or "Aptitude Trainer"
    row["Company Mail"] = clean_text(row.get("Company Mail"))
    row["Phone Number"] = clean_text(row.get("Phone Number"))
    row["Date of Joining"] = format_date_value(row.get("Date of Joining"))
    row["Today Date"] = format_date_value(row.get("Today Date"))
    row["Qualification"] = clean_text(row.get("Qualification"))
    row["Certifications"] = clean_text(row.get("Certifications"))
    row["Skills"] = clean_text(row.get("Skills"))
    existing_exp = clean_text(row.get("Experience"))
    row["Experience"] = existing_exp or calculate_experience(row.get("Date of Joining"), row.get("Today Date"))
    return row


def read_rows(excel_path):
    workbook = load_workbook(excel_path, data_only=True)
    worksheet, header_row = find_data_sheet(workbook)
    headers = [normalize_header(cell.value) for cell in worksheet[header_row]]
    rows = []
    for excel_row in range(header_row + 1, worksheet.max_row + 1):
        values = {}
        has_content = False
        for column_index, header in enumerate(headers, start=1):
            if not header:
                continue
            value = worksheet.cell(row=excel_row, column=column_index).value
            if value not in [None, ""]:
                has_content = True
            values[header] = value
        if not has_content:
            continue
        if not clean_text(values.get("Name")) and not clean_text(values.get("Employee ID")):
            continue
        values["_excel_row"] = excel_row
        rows.append(prepare_row(values))
    return workbook, worksheet, rows


# =========================================================
# EMBEDDED EXCEL PHOTO SUPPORT
# =========================================================


def image_anchor_row(image):
    try:
        return image.anchor._from.row + 1
    except Exception:
        return None


def save_embedded_image(image, output_path):
    try:
        image_bytes = image._data()
        pil_image = PILImage.open(BytesIO(image_bytes))
        if pil_image.mode == "RGBA":
            background = PILImage.new("RGB", pil_image.size, "white")
            background.paste(pil_image, mask=pil_image.getchannel("A"))
            pil_image = background
        elif pil_image.mode != "RGB":
            pil_image = pil_image.convert("RGB")
        pil_image.save(output_path, format="PNG")
        return str(output_path)
    except Exception:
        return None


def extract_embedded_photos(worksheet):
    photo_map = {}
    for index, image in enumerate(getattr(worksheet, "_images", []), start=1):
        excel_row = image_anchor_row(image)
        if excel_row is None:
            continue
        output_path = TEMP_IMAGE_DIR / f"excel_row_{excel_row}_{index}.png"
        saved = save_embedded_image(image, output_path)
        if saved:
            photo_map[excel_row] = saved
    return photo_map


def resolve_photo_from_cell(photo_value):
    photo_value = clean_text(photo_value)
    if not photo_value:
        return None
    candidate = Path(photo_value)
    candidates = [candidate] if candidate.is_absolute() else [
        ROOT_DIR / photo_value,
        ROOT_DIR / "data" / photo_value,
        ROOT_DIR / "data" / "CTImages" / photo_value,
        ROOT_DIR / "data" / "Aptitude_Images" / photo_value,
        ROOT_DIR / "images" / photo_value,
    ]
    for path in candidates:
        if path.exists() and path.is_file():
            return str(path)
    return None


# =========================================================
# WORD TEMPLATE HELPERS
# =========================================================


def iter_paragraphs(parent):
    for paragraph in parent.paragraphs:
        yield paragraph
    for table in parent.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from iter_paragraphs(cell)


def paragraph_text(paragraph):
    return "".join(run.text for run in paragraph.runs)


def replace_in_paragraph(paragraph, replacements):
    for run in paragraph.runs:
        for old, new in replacements.items():
            if old in run.text:
                run.text = run.text.replace(old, clean_text(new))
    full_text = paragraph_text(paragraph)
    if not any(old in full_text for old in replacements):
        return
    new_text = full_text
    for old, new in replacements.items():
        new_text = new_text.replace(old, clean_text(new))
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(new_text)


def replace_all_placeholders(document, replacements):
    for paragraph in iter_paragraphs(document):
        replace_in_paragraph(paragraph, replacements)
    for section in document.sections:
        for paragraph in iter_paragraphs(section.header):
            replace_in_paragraph(paragraph, replacements)
        for paragraph in iter_paragraphs(section.footer):
            replace_in_paragraph(paragraph, replacements)


def insert_photo(document, photo_path):
    for paragraph in iter_paragraphs(document):
        full_text = paragraph_text(paragraph)
        placeholder = next((p for p in ["{{Image}}", "{{Photo}}"] if p in full_text), None)
        if placeholder is None:
            continue
        for run in paragraph.runs:
            run.text = run.text.replace(placeholder, "")
        if photo_path and Path(photo_path).exists():
            picture_run = paragraph.add_run()
            picture_run.add_picture(str(photo_path), width=Inches(1.45))
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return


def remove_remaining_placeholders(document):
    pattern = re.compile(r"\{\{[^{}]+\}\}")
    containers = [document]
    for section in document.sections:
        containers.extend([section.header, section.footer])
    for container in containers:
        for paragraph in iter_paragraphs(container):
            full_text = paragraph_text(paragraph)
            cleaned = pattern.sub("", full_text)
            if cleaned == full_text:
                continue
            if paragraph.runs:
                paragraph.runs[0].text = cleaned
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.add_run(cleaned)


def build_context(row):
    skills = split_skills(row.get("Skills"))
    exam_text = build_exam_expertise(skills)
    return {
        "{{Name}}": row.get("Name", ""),
        "{{Emp ID}}": row.get("Employee ID", ""),
        "{{Employee ID}}": row.get("Employee ID", ""),
        "{{Designation}}": row.get("Designation", "Aptitude Trainer"),
        "{{Company Mail}}": row.get("Company Mail", ""),
        "{{Phone Number}}": row.get("Phone Number", ""),
        "{{Date of Joining}}": row.get("Date of Joining", ""),
        "{{Today Date}}": row.get("Today Date", "") or date.today().strftime("%d/%m/%Y"),
        "{{Experience}}": row.get("Experience", ""),
        "{{Qualification}}": row.get("Qualification", ""),
        "{{Certifications}}": build_certifications(row.get("Certifications")),
        "{{Skills}}": format_core_skills(skills),
        "{{Summary}}": build_professional_summary(row, skills),
        "{{Professional Highlights}}": build_professional_highlights(skills),
        "{{Subjects Handled}}": build_subjects_handled(skills),
        "{{Exam Placement Expertise}}": exam_text,
        "{{Exam / Placement Expertise}}": exam_text,
        "{{Training Expertise}}": build_training_expertise(skills),
        "{{Core Competencies}}": build_core_competencies(skills),
        "{{Languages}}": "",
        "{{Rating}}": "",
        "{{Training Projects}}": "Project allocation not yet done.",
        "Profile - Technical Trainer": "Profile - Aptitude Trainer",
    }


def generate_word_profile(row, photo_path):
    document = Document(str(TEMPLATE_PATH))
    insert_photo(document, photo_path)
    replace_all_placeholders(document, build_context(row))
    remove_remaining_placeholders(document)
    employee_id = safe_filename(row.get("Employee ID"))
    name = safe_filename(row.get("Name"))
    base_name = f"{employee_id}_{name}" if row.get("Employee ID") else name
    output_path = WORD_OUTPUT_DIR / f"{base_name}.docx"
    document.save(output_path)
    return output_path


# =========================================================
# PDF SUPPORT
# =========================================================


def get_pdf_conversion_status():
    if os.name == "nt":
        try:
            import docx2pdf  # noqa: F401
            return True, "Microsoft Word / docx2pdf"
        except Exception:
            return False, "docx2pdf is not available"
    libreoffice = shutil.which("libreoffice") or shutil.which("soffice")
    return (True, "LibreOffice") if libreoffice else (False, "LibreOffice is not available")


def convert_docx_to_pdf(docx_path):
    docx_path = Path(docx_path).resolve()
    pdf_path = (PDF_OUTPUT_DIR / f"{docx_path.stem}.pdf").resolve()
    if os.name == "nt":
        from docx2pdf import convert
        convert(str(docx_path), str(pdf_path))
        if not pdf_path.exists():
            raise RuntimeError("PDF was not created. Microsoft Word must be installed.")
        return pdf_path
    libreoffice = shutil.which("libreoffice") or shutil.which("soffice")
    if not libreoffice:
        raise RuntimeError("LibreOffice is not installed.")
    process = subprocess.run(
        [libreoffice, "--headless", "--convert-to", "pdf", "--outdir", str(PDF_OUTPUT_DIR), str(docx_path)],
        capture_output=True,
        text=True,
        timeout=120,
    )
    if process.returncode != 0:
        raise RuntimeError(process.stderr or process.stdout or "LibreOffice PDF conversion failed.")
    if not pdf_path.exists():
        raise RuntimeError("LibreOffice completed but PDF was not found.")
    return pdf_path


def create_zip(generated_files, output_type):
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as zip_file:
        for item in generated_files:
            word_path = item.get("word")
            pdf_path = item.get("pdf")
            if output_type in ["Word", "Word + PDF"] and word_path and Path(word_path).exists():
                zip_file.write(word_path, arcname="Word/" + Path(word_path).name)
            if output_type in ["PDF", "Word + PDF"] and pdf_path and Path(pdf_path).exists():
                zip_file.write(pdf_path, arcname="PDF/" + Path(pdf_path).name)
    return buffer.getvalue()


# =========================================================
# PAGE
# =========================================================

st.markdown(
    '<h1 style="color:#1F4E79;margin-bottom:0;">🧠 Aptitude Trainer Profile Generator</h1>',
    unsafe_allow_html=True,
)
st.caption("Independent Aptitude Trainer profile generation from data/Aptitude_Trainer.xlsx")
st.info(
    "This page uses only Aptitude_Trainer.xlsx and Aptitude_Trainer_Profile_Template.docx. "
    "It does not modify Trainers.xlsx or the existing Technical Trainer profile generator."
)

st.markdown("---")
col1, col2 = st.columns(2)
with col1:
    if EXCEL_PATH.exists():
        st.success("Excel found: " + EXCEL_PATH.name)
    else:
        st.error("Excel file missing: " + str(EXCEL_PATH))
with col2:
    if TEMPLATE_PATH.exists():
        st.success("Template found: " + TEMPLATE_PATH.name)
    else:
        st.error("Template missing: " + str(TEMPLATE_PATH))

if not EXCEL_PATH.exists() or not TEMPLATE_PATH.exists():
    st.stop()

try:
    workbook, worksheet, rows = read_rows(EXCEL_PATH)
    embedded_photo_map = extract_embedded_photos(worksheet)
except Exception as error:
    st.error("Unable to read Aptitude_Trainer.xlsx")
    st.exception(error)
    st.stop()

if not rows:
    st.warning("No Aptitude Trainer records were found.")
    st.stop()

st.markdown("---")
st.subheader("Aptitude Trainers")
search_text = st.text_input("Search by Employee ID or Name", placeholder="Example: CT0567 or Jyothsna")

filtered_rows = []
for row in rows:
    source = (clean_text(row.get("Employee ID")) + " " + clean_text(row.get("Name"))).lower()
    if not search_text.strip() or search_text.strip().lower() in source:
        filtered_rows.append(row)

st.dataframe(
    [
        {
            "Employee ID": row.get("Employee ID", ""),
            "Name": row.get("Name", ""),
            "Designation": row.get("Designation", ""),
            "Qualification": row.get("Qualification", ""),
            "Experience": row.get("Experience", ""),
        }
        for row in filtered_rows
    ],
    hide_index=True,
    use_container_width=True,
)
st.caption(f"{len(filtered_rows)} of {len(rows)} trainer(s) shown.")

option_to_row = {}
options = []
for row in filtered_rows:
    option = (clean_text(row.get("Employee ID")) + " - " + clean_text(row.get("Name"))).strip(" -")
    options.append(option)
    option_to_row[option] = row

selection_mode = st.radio("Generate Profiles For", ["All Trainers", "Selected Trainers"], horizontal=True)
if selection_mode == "Selected Trainers":
    selected_options = st.multiselect("Select Trainers", options)
    selected_rows = [option_to_row[option] for option in selected_options]
else:
    selected_rows = filtered_rows

st.markdown("---")
output_type = st.radio("Output Format", ["Word", "PDF", "Word + PDF"], horizontal=True)

if output_type in ["PDF", "Word + PDF"]:
    pdf_available, pdf_method = get_pdf_conversion_status()
    if pdf_available:
        st.success("PDF conversion available: " + pdf_method)
    else:
        st.warning("PDF conversion unavailable: " + pdf_method)

button_label = f"🚀 Generate {len(selected_rows)} Aptitude Trainer Profile" + ("s" if len(selected_rows) != 1 else "")
if st.button(button_label, type="primary", use_container_width=True, disabled=(len(selected_rows) == 0)):
    generated_files = []
    errors = []
    progress = st.progress(0)
    total = len(selected_rows)

    for position, source_row in enumerate(selected_rows, start=1):
        try:
            row = prepare_row(source_row)
            excel_row = row.get("_excel_row")
            photo_path = embedded_photo_map.get(excel_row) or resolve_photo_from_cell(row.get("Photo"))
            word_path = generate_word_profile(row, photo_path)
            pdf_path = None
            if output_type in ["PDF", "Word + PDF"]:
                pdf_path = convert_docx_to_pdf(word_path)
            generated_files.append(
                {
                    "Employee ID": row.get("Employee ID", ""),
                    "Name": row.get("Name", ""),
                    "word": word_path,
                    "pdf": pdf_path,
                }
            )
        except Exception as error:
            errors.append(
                {
                    "Employee ID": source_row.get("Employee ID", ""),
                    "Name": source_row.get("Name", ""),
                    "Error": str(error),
                }
            )
        progress.progress(position / total)

    progress.empty()
    st.session_state["aptitude_generated_files"] = generated_files
    st.session_state["aptitude_output_type"] = output_type
    st.session_state["aptitude_zip"] = create_zip(generated_files, output_type) if generated_files else None

    if generated_files:
        st.success(f"{len(generated_files)} aptitude trainer profile(s) generated successfully.")
    if errors:
        st.warning(f"{len(errors)} profile(s) could not be generated.")
        st.dataframe(errors, hide_index=True, use_container_width=True)


generated_files = st.session_state.get("aptitude_generated_files", [])
zip_bytes = st.session_state.get("aptitude_zip")
saved_output_type = st.session_state.get("aptitude_output_type", "Word")

if generated_files:
    st.markdown("---")
    st.subheader("Download Generated Profiles")
    st.dataframe(
        [
            {
                "Employee ID": item.get("Employee ID", ""),
                "Name": item.get("Name", ""),
                "Word": "Generated" if item.get("word") else "",
                "PDF": "Generated" if item.get("pdf") else "",
            }
            for item in generated_files
        ],
        hide_index=True,
        use_container_width=True,
    )

    if zip_bytes:
        st.download_button(
            "⬇️ Download All Profiles as ZIP",
            data=zip_bytes,
            file_name="Aptitude_Trainer_Profiles.zip",
            mime="application/zip",
            use_container_width=True,
        )

    with st.expander("Individual Downloads"):
        for index, item in enumerate(generated_files):
            st.markdown(f"**{item.get('Employee ID', '')} - {item.get('Name', '')}**")
            word_col, pdf_col = st.columns(2)
            word_path = item.get("word")
            if saved_output_type in ["Word", "Word + PDF"] and word_path and Path(word_path).exists():
                with word_col:
                    st.download_button(
                        "⬇️ Word",
                        data=Path(word_path).read_bytes(),
                        file_name=Path(word_path).name,
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        key=f"apt_word_{index}",
                        use_container_width=True,
                    )
            pdf_path = item.get("pdf")
            if saved_output_type in ["PDF", "Word + PDF"] and pdf_path and Path(pdf_path).exists():
                with pdf_col:
                    st.download_button(
                        "⬇️ PDF",
                        data=Path(pdf_path).read_bytes(),
                        file_name=Path(pdf_path).name,
                        mime="application/pdf",
                        key=f"apt_pdf_{index}",
                        use_container_width=True,
                    )
