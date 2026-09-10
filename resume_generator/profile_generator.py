from pathlib import Path
from copy import deepcopy
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# =========================================================
# ITERATE THROUGH ALL PARAGRAPHS
# INCLUDING TABLE CELLS
# =========================================================

def iter_paragraphs(parent):
    """
    Iterate through all paragraphs, including
    paragraphs inside table cells.
    """

    for paragraph in parent.paragraphs:
        yield paragraph

    for table in parent.tables:

        for row in table.rows:

            for cell in row.cells:

                yield from iter_paragraphs(cell)


# =========================================================
# GET FULL PARAGRAPH TEXT
# =========================================================

def get_paragraph_text(paragraph):
    """
    Get complete paragraph text even when Word
    has divided the text into multiple runs.
    """

    return "".join(
        run.text
        for run in paragraph.runs
    )


# =========================================================
# REPLACE TEXT IN PARAGRAPH
# =========================================================

def replace_text_in_paragraph(
    paragraph,
    replacements
):
    """
    Replace placeholders while preserving
    the existing Word formatting as much as possible.
    """

    # -----------------------------------------------------
    # First try replacing inside individual runs
    # -----------------------------------------------------

    for run in paragraph.runs:

        for old_text, new_text in replacements.items():

            if old_text in run.text:

                run.text = run.text.replace(
                    old_text,
                    str(new_text or "")
                )

    # -----------------------------------------------------
    # Word can split a placeholder across runs.
    # Handle that case also.
    # -----------------------------------------------------

    full_text = get_paragraph_text(
        paragraph
    )

    found = False

    for old_text in replacements:

        if old_text in full_text:

            found = True
            break

    if not found:
        return

    new_full_text = full_text

    for old_text, new_text in replacements.items():

        new_full_text = (
            new_full_text.replace(
                old_text,
                str(new_text or "")
            )
        )

    if paragraph.runs:

        paragraph.runs[0].text = (
            new_full_text
        )

        for run in paragraph.runs[1:]:

            run.text = ""

    else:

        paragraph.add_run(
            new_full_text
        )


# =========================================================
# REPLACE ALL PLACEHOLDERS
# =========================================================

def replace_all_placeholders(
    document,
    replacements
):
    """
    Replace placeholders in body, tables,
    headers and footers.
    """

    # -----------------------------------------------------
    # Document body
    # -----------------------------------------------------

    for paragraph in iter_paragraphs(
        document
    ):

        replace_text_in_paragraph(
            paragraph,
            replacements
        )

    # -----------------------------------------------------
    # Headers and footers
    # -----------------------------------------------------

    for section in document.sections:

        for paragraph in iter_paragraphs(
            section.header
        ):

            replace_text_in_paragraph(
                paragraph,
                replacements
            )

        for paragraph in iter_paragraphs(
            section.footer
        ):

            replace_text_in_paragraph(
                paragraph,
                replacements
            )


# =========================================================
# INSERT PHOTO
# =========================================================

def insert_photo(
    document,
    photo_path
):
    """
    Find {{Image}} and replace it with
    the extracted candidate photo.
    """

    if not photo_path:
        return False

    photo_path = Path(
        photo_path
    )

    if not photo_path.exists():
        return False

    for paragraph in iter_paragraphs(
        document
    ):

        full_text = get_paragraph_text(
            paragraph
        )

        if "{{Image}}" not in full_text:
            continue

        # -------------------------------------------------
        # Remove image placeholder
        # -------------------------------------------------

        for run in paragraph.runs:

            run.text = run.text.replace(
                "{{Image}}",
                ""
            )

        # -------------------------------------------------
        # Insert candidate photo
        # -------------------------------------------------

        picture_run = (
            paragraph.add_run()
        )

        picture_run.add_picture(
            str(photo_path),
            width=Inches(1.45)
        )

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        return True

    return False


# =========================================================
# TABLE CELL BACKGROUND
# =========================================================

def set_cell_shading(
    cell,
    fill
):
    """
    Set table cell background color.
    """

    tc_pr = (
        cell._tc.get_or_add_tcPr()
    )

    shading = OxmlElement(
        "w:shd"
    )

    shading.set(
        qn("w:fill"),
        fill
    )

    tc_pr.append(
        shading
    )


# =========================================================
# TABLE CELL VERTICAL ALIGNMENT
# =========================================================

def set_cell_vertical_center(
    cell
):
    """
    Vertically center table cell contents.
    """

    tc_pr = (
        cell._tc.get_or_add_tcPr()
    )

    valign = OxmlElement(
        "w:vAlign"
    )

    valign.set(
        qn("w:val"),
        "center"
    )

    tc_pr.append(
        valign
    )


# =========================================================
# TABLE BORDERS
# =========================================================

def set_table_borders(
    table
):
    """
    Add table borders directly through Word XML.

    Important:
    This does NOT depend on the 'Table Grid'
    Word style.
    """

    table_xml = (
        table._tbl
    )

    table_properties = (
        table_xml.tblPr
    )

    borders = (
        table_properties.first_child_found_in(
            "w:tblBorders"
        )
    )

    if borders is None:

        borders = OxmlElement(
            "w:tblBorders"
        )

        table_properties.append(
            borders
        )

    border_names = [
        "top",
        "left",
        "bottom",
        "right",
        "insideH",
        "insideV"
    ]

    for border_name in border_names:

        tag_name = (
            "w:" + border_name
        )

        border = borders.find(
            qn(tag_name)
        )

        if border is None:

            border = OxmlElement(
                tag_name
            )

            borders.append(
                border
            )

        border.set(
            qn("w:val"),
            "single"
        )

        border.set(
            qn("w:sz"),
            "6"
        )

        border.set(
            qn("w:space"),
            "0"
        )

        border.set(
            qn("w:color"),
            "808080"
        )


# =========================================================
# FIND PARAGRAPH
# =========================================================

def find_paragraph(
    document,
    text
):
    """
    Find a paragraph by visible text.
    """

    target = (
        text.strip()
        .lower()
        .rstrip(":")
    )

    for paragraph in iter_paragraphs(
        document
    ):

        current = (
            get_paragraph_text(
                paragraph
            )
            .strip()
            .lower()
            .rstrip(":")
        )

        if current == target:
            return paragraph

    return None


# =========================================================
# CREATE SECTION HEADING
# =========================================================

def create_section_heading(
    document,
    heading_text,
    reference_heading="Professional Certifications"
):
    """
    Create a new section heading.

    Whenever possible, formatting is copied from
    an existing template heading so that the new
    Experience section matches the original design.
    """

    new_paragraph = (
        document.add_paragraph()
    )

    reference = find_paragraph(
        document,
        reference_heading
    )

    # -----------------------------------------------------
    # Copy paragraph formatting from existing heading
    # -----------------------------------------------------

    if (
        reference is not None
        and reference._p.pPr is not None
    ):

        new_paragraph._p.insert(
            0,
            deepcopy(
                reference._p.pPr
            )
        )

    new_run = (
        new_paragraph.add_run(
            heading_text
        )
    )

    # -----------------------------------------------------
    # Copy run formatting from reference heading
    # -----------------------------------------------------

    if reference is not None:

        for reference_run in (
            reference.runs
        ):

            if reference_run.text.strip():

                if (
                    reference_run._r.rPr
                    is not None
                ):

                    new_run._r.insert(
                        0,
                        deepcopy(
                            reference_run._r.rPr
                        )
                    )

                break

    else:

        # -------------------------------------------------
        # Fallback formatting
        # -------------------------------------------------

        new_run.bold = True

        new_run.font.size = (
            Pt(12)
        )

        new_run.font.color.rgb = (
            RGBColor(
                31,
                78,
                121
            )
        )

    return new_paragraph


# =========================================================
# EXPERIENCE TABLE WIDTHS
# =========================================================

def set_experience_table_widths(
    table
):
    """
    Set approximate column widths for:

    S.No
    Name of Organization
    Years Worked
    """

    widths = [
        Inches(0.60),
        Inches(4.55),
        Inches(2.10)
    ]

    for row in table.rows:

        for index, width in enumerate(
            widths
        ):

            if index < len(
                row.cells
            ):

                row.cells[
                    index
                ].width = width


# =========================================================
# CLEAN EXPERIENCE VALUE
# =========================================================

def clean_experience_value(
    value
):

    if value is None:
        return ""

    value = str(
        value
    ).strip()

    if value.lower() in [
        "none",
        "nan"
    ]:

        return ""

    return value


# =========================================================
# ADD PROFESSIONAL EXPERIENCE TABLE
# =========================================================

def add_experience_table(
    document,
    experience_details
):
    """
    Add professional experience only when
    organization + employment period information
    has been extracted from the resume.

    Format:

    S.No | Name of Organization | Years Worked
    """

    if not experience_details:
        return

    valid_experience = []

    for item in experience_details:

        organization = (
            clean_experience_value(
                item.get(
                    "Name of Organization",
                    ""
                )
            )
        )

        years_worked = (
            clean_experience_value(
                item.get(
                    "Years Worked",
                    ""
                )
            )
        )

        # -------------------------------------------------
        # Organization is required.
        # Do not invent organization details.
        # -------------------------------------------------

        if not organization:
            continue

        valid_experience.append(
            {
                "Name of Organization":
                    organization,

                "Years Worked":
                    years_worked
            }
        )

    if not valid_experience:
        return

    # -----------------------------------------------------
    # We insert the table immediately before
    # Professional Certifications.
    # -----------------------------------------------------

    certifications_heading = (
        find_paragraph(
            document,
            "Professional Certifications"
        )
    )

    # -----------------------------------------------------
    # Fallback if template heading differs
    # -----------------------------------------------------

    if certifications_heading is None:

        certifications_heading = (
            find_paragraph(
                document,
                "Professional Highlights"
            )
        )

    if certifications_heading is None:
        return

    # -----------------------------------------------------
    # Create Experience heading
    # -----------------------------------------------------

    experience_heading = (
        create_section_heading(
            document,
            "Professional Experience"
        )
    )

    # -----------------------------------------------------
    # Create table
    # -----------------------------------------------------

    table = document.add_table(
        rows=1,
        cols=3
    )

    table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )

    set_table_borders(
        table
    )

    # -----------------------------------------------------
    # Header row
    # -----------------------------------------------------

    headers = [
        "S.No",
        "Name of Organization",
        "Years Worked"
    ]

    header_cells = (
        table.rows[0].cells
    )

    for index, header in enumerate(
        headers
    ):

        cell = (
            header_cells[index]
        )

        cell.text = header

        set_cell_shading(
            cell,
            "1F4E78"
        )

        set_cell_vertical_center(
            cell
        )

        for paragraph in (
            cell.paragraphs
        ):

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            for run in (
                paragraph.runs
            ):

                run.bold = True

                run.font.size = (
                    Pt(9)
                )

                run.font.color.rgb = (
                    RGBColor(
                        255,
                        255,
                        255
                    )
                )

    # -----------------------------------------------------
    # Experience rows
    # -----------------------------------------------------

    for serial_no, item in enumerate(
        valid_experience,
        start=1
    ):

        cells = (
            table.add_row().cells
        )

        cells[0].text = str(
            serial_no
        )

        cells[1].text = (
            item[
                "Name of Organization"
            ]
        )

        cells[2].text = (
            item[
                "Years Worked"
            ]
        )

        for cell_index, cell in enumerate(
            cells
        ):

            set_cell_vertical_center(
                cell
            )

            for paragraph in (
                cell.paragraphs
            ):

                if cell_index == 0:

                    paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.CENTER
                    )

                else:

                    paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.LEFT
                    )

                for run in (
                    paragraph.runs
                ):

                    run.font.size = (
                        Pt(9)
                    )

    # -----------------------------------------------------
    # Apply borders again after adding rows
    # -----------------------------------------------------

    set_table_borders(
        table
    )

    set_experience_table_widths(
        table
    )

    # -----------------------------------------------------
    # Position:
    #
    # Professional Summary
    # Summary content
    # Professional Experience
    # Experience table
    # Professional Certifications
    # -----------------------------------------------------

    certifications_heading._p.addprevious(
        experience_heading._p
    )

    certifications_heading._p.addprevious(
        table._tbl
    )


# =========================================================
# PROJECT TABLE WIDTHS
# =========================================================

def set_project_table_widths(
    table
):
    """
    Set approximate widths for project table.
    """

    widths = [
        Inches(0.55),
        Inches(3.10),
        Inches(1.35),
        Inches(2.25)
    ]

    for row in table.rows:

        for index, width in enumerate(
            widths
        ):

            if index < len(
                row.cells
            ):

                row.cells[
                    index
                ].width = width


# =========================================================
# FIND TRAINING PROJECTS HEADING
# =========================================================

def find_training_projects_heading(
    document
):
    """
    Locate Training Projects heading.
    """

    for paragraph in iter_paragraphs(
        document
    ):

        text = (
            get_paragraph_text(
                paragraph
            )
            .strip()
            .lower()
        )

        if text.startswith(
            "training projects"
        ):

            return paragraph

    return None


# =========================================================
# REMOVE PROJECT PLACEHOLDER
# =========================================================

def remove_empty_project_placeholder(
    document
):
    """
    Remove {{Training Projects}} if it exists.
    """

    replace_all_placeholders(
        document,
        {
            "{{Training Projects}}": ""
        }
    )


# =========================================================
# CLEAN PROJECT VALUE
# =========================================================

def clean_project_value(
    value
):

    if value is None:
        return ""

    value = str(
        value
    ).strip()

    if value.lower() in [
        "none",
        "nan"
    ]:

        return ""

    return value


# =========================================================
# ADD TRAINING PROJECTS
# =========================================================

def add_training_projects(
    document,
    projects
):
    """
    Create:

    S.No
    College / Client
    Location
    Domain / Subject Area
    """

    heading = (
        find_training_projects_heading(
            document
        )
    )

    if heading is None:
        return

    valid_projects = []

    if projects:

        for project in projects:

            college = (
                clean_project_value(
                    project.get(
                        "College / Client",
                        ""
                    )
                )
            )

            location = (
                clean_project_value(
                    project.get(
                        "Location",
                        ""
                    )
                )
            )

            subject = (
                clean_project_value(
                    project.get(
                        "Domain / Subject Area",
                        ""
                    )
                )
            )

            if (
                college
                or location
                or subject
            ):

                valid_projects.append(
                    {
                        "College / Client":
                            college,

                        "Location":
                            location,

                        "Domain / Subject Area":
                            subject
                    }
                )

    # -----------------------------------------------------
    # No projects
    # -----------------------------------------------------

    if not valid_projects:

        no_project_para = (
            document.add_paragraph()
        )

        no_project_run = (
            no_project_para.add_run(
                "Project allocation not yet done."
            )
        )

        no_project_run.font.size = (
            Pt(10)
        )

        heading._p.addnext(
            no_project_para._p
        )

        return

    # -----------------------------------------------------
    # Completed Projects heading
    # -----------------------------------------------------

    completed_para = (
        document.add_paragraph()
    )

    completed_run = (
        completed_para.add_run(
            "Completed Projects:"
        )
    )

    completed_run.bold = True

    completed_run.font.size = (
        Pt(11)
    )

    # -----------------------------------------------------
    # Project table
    # -----------------------------------------------------

    table = document.add_table(
        rows=1,
        cols=4
    )

    table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )

    set_table_borders(
        table
    )

    headers = [
        "S.No",
        "College / Client",
        "Location",
        "Domain / Subject Area"
    ]

    header_cells = (
        table.rows[0].cells
    )

    for index, header in enumerate(
        headers
    ):

        cell = (
            header_cells[index]
        )

        cell.text = header

        set_cell_shading(
            cell,
            "1F4E78"
        )

        set_cell_vertical_center(
            cell
        )

        for paragraph in (
            cell.paragraphs
        ):

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            for run in (
                paragraph.runs
            ):

                run.bold = True

                run.font.size = (
                    Pt(9)
                )

                run.font.color.rgb = (
                    RGBColor(
                        255,
                        255,
                        255
                    )
                )

    # -----------------------------------------------------
    # Project rows
    # -----------------------------------------------------

    for serial_no, project in enumerate(
        valid_projects,
        start=1
    ):

        cells = (
            table.add_row().cells
        )

        cells[0].text = str(
            serial_no
        )

        cells[1].text = (
            project[
                "College / Client"
            ]
        )

        cells[2].text = (
            project[
                "Location"
            ]
        )

        cells[3].text = (
            project[
                "Domain / Subject Area"
            ]
        )

        for cell_index, cell in enumerate(
            cells
        ):

            set_cell_vertical_center(
                cell
            )

            for paragraph in (
                cell.paragraphs
            ):

                if cell_index == 0:

                    paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.CENTER
                    )

                else:

                    paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.LEFT
                    )

                for run in (
                    paragraph.runs
                ):

                    run.font.size = (
                        Pt(8.5)
                    )

    set_table_borders(
        table
    )

    set_project_table_widths(
        table
    )

    heading._p.addnext(
        completed_para._p
    )

    completed_para._p.addnext(
        table._tbl
    )


# =========================================================
# REMOVE REMAINING PLACEHOLDERS
# =========================================================

def remove_remaining_placeholders(
    document
):
    """
    Remove unused {{...}} placeholders.
    """

    placeholder_pattern = (
        re.compile(
            r"\{\{[^{}]+\}\}"
        )
    )

    for paragraph in iter_paragraphs(
        document
    ):

        full_text = (
            get_paragraph_text(
                paragraph
            )
        )

        if not full_text:
            continue

        cleaned_text = (
            placeholder_pattern.sub(
                "",
                full_text
            )
        )

        if cleaned_text == full_text:
            continue

        if paragraph.runs:

            paragraph.runs[0].text = (
                cleaned_text
            )

            for run in (
                paragraph.runs[1:]
            ):

                run.text = ""

        else:

            paragraph.add_run(
                cleaned_text
            )


# =========================================================
# REMOVE HEADER / FOOTER PLACEHOLDERS
# =========================================================

def remove_header_footer_placeholders(
    document
):

    placeholder_pattern = (
        re.compile(
            r"\{\{[^{}]+\}\}"
        )
    )

    for section in document.sections:

        for container in [
            section.header,
            section.footer
        ]:

            for paragraph in iter_paragraphs(
                container
            ):

                full_text = (
                    get_paragraph_text(
                        paragraph
                    )
                )

                cleaned_text = (
                    placeholder_pattern.sub(
                        "",
                        full_text
                    )
                )

                if cleaned_text == full_text:
                    continue

                if paragraph.runs:

                    paragraph.runs[
                        0
                    ].text = (
                        cleaned_text
                    )

                    for run in (
                        paragraph.runs[1:]
                    ):

                        run.text = ""

                else:

                    paragraph.add_run(
                        cleaned_text
                    )


# =========================================================
# MAIN PROFILE GENERATOR
# =========================================================

def generate_profile(
    template_path,
    output_path,
    data,
    profile_type,
    photo_path=None,
    projects=None,
    experience_details=None
):
    """
    Generate Technical Trainer or Aptitude Trainer
    profile using the selected Word template.
    """

    template_path = Path(
        template_path
    )

    output_path = Path(
        output_path
    )

    # -----------------------------------------------------
    # Validate template
    # -----------------------------------------------------

    if not template_path.exists():

        raise FileNotFoundError(
            f"Template not found: "
            f"{template_path}"
        )

    # -----------------------------------------------------
    # Validate profile type
    # -----------------------------------------------------

    valid_profile_types = [
        "Technical Trainer",
        "Aptitude Trainer"
    ]

    if profile_type not in (
        valid_profile_types
    ):

        raise ValueError(
            "Profile type must be "
            "'Technical Trainer' or "
            "'Aptitude Trainer'."
        )

    # -----------------------------------------------------
    # Create output directory
    # -----------------------------------------------------

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    # -----------------------------------------------------
    # Open the selected template.
    #
    # Logo and existing design are retained.
    # -----------------------------------------------------

    document = Document(
        str(template_path)
    )

    # =====================================================
    # PHOTO
    # =====================================================

    insert_photo(
        document,
        photo_path
    )

    # =====================================================
    # EXPERIENCE TABLE
    #
    # Insert BEFORE normal placeholder replacement so
    # template section formatting can be copied.
    # =====================================================

    if experience_details is None:

        experience_details = (
            data.get(
                "Experience Details",
                []
            )
        )

    add_experience_table(
        document,
        experience_details
    )

    # =====================================================
    # PLACEHOLDER VALUES
    # =====================================================

    replacements = {

        "{{Name}}":
            data.get(
                "Name",
                ""
            ),

        "{{Emp ID}}":
            data.get(
                "Emp ID",
                ""
            ),

        "{{Designation}}":
            data.get(
                "Designation",
                profile_type
            ),

        "{{Qualification}}":
            data.get(
                "Qualification",
                ""
            ),

        "{{Date of Joining}}":
            data.get(
                "Date of Joining",
                ""
            ),

        "{{Experience}}":
            data.get(
                "Experience",
                ""
            ),

        "{{Company Mail}}":
            data.get(
                "Company Mail",
                ""
            ),

        "{{Phone Number}}":
            data.get(
                "Phone Number",
                ""
            ),

        "{{Rating}}":
            data.get(
                "Rating",
                ""
            ),

        "{{Skills}}":
            data.get(
                "Skills",
                ""
            ),

        "{{Summary}}":
            data.get(
                "Summary",
                ""
            ),

        "{{Certifications}}":
            data.get(
                "Certifications",
                ""
            ),

        "{{Professional Highlights}}":
            data.get(
                "Professional Highlights",
                ""
            ),

        "{{Training Expertise}}":
            data.get(
                "Training Expertise",
                ""
            ),

        "{{Core Competencies}}":
            data.get(
                "Core Competencies",
                ""
            ),

        # -------------------------------------------------
        # Aptitude Trainer
        # -------------------------------------------------

        "{{Subjects Handled}}":
            data.get(
                "Subjects Handled",
                ""
            ),

        "{{Exam Placement Expertise}}":
            data.get(
                "Exam Placement Expertise",
                ""
            ),

        "{{Exam / Placement Expertise}}":
            data.get(
                "Exam Placement Expertise",
                ""
            ),

        "{{Languages}}":
            data.get(
                "Languages",
                ""
            ),

        # -------------------------------------------------
        # Technical Trainer
        # -------------------------------------------------

        "{{LeetCode Profile Link}}":
            data.get(
                "LeetCode Profile Link",
                ""
            ),

        "{{No of Problems Done}}":
            data.get(
                "No of Problems Done",
                ""
            ),

        "{{Other Profiles}}":
            data.get(
                "Other Profiles",
                ""
            )
    }

    # =====================================================
    # PROFILE TITLE
    # =====================================================

    if profile_type == (
        "Aptitude Trainer"
    ):

        replacements[
            "Profile - Technical Trainer"
        ] = (
            "Profile - Aptitude Trainer"
        )

    else:

        replacements[
            "Profile - Aptitude Trainer"
        ] = (
            "Profile - Technical Trainer"
        )

    # =====================================================
    # REPLACE PLACEHOLDERS
    # =====================================================

    replace_all_placeholders(
        document,
        replacements
    )

    # =====================================================
    # TRAINING PROJECTS
    # =====================================================

    remove_empty_project_placeholder(
        document
    )

    if projects is None:

        projects = (
            data.get(
                "Training Projects",
                []
            )
        )

    add_training_projects(
        document,
        projects
    )

    # =====================================================
    # CLEAN UNUSED PLACEHOLDERS
    # =====================================================

    remove_remaining_placeholders(
        document
    )

    remove_header_footer_placeholders(
        document
    )

    # =====================================================
    # SAVE WORD FILE
    # =====================================================

    document.save(
        str(output_path)
    )

    return str(
        output_path
    )