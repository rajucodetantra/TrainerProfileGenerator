import streamlit as st

from login import require_login


# =========================================================
# PAGE CONFIG
# =========================================================

st.set_page_config(
    page_title="Page Name",
    page_icon="📄",
    layout="wide"
)


# =========================================================
# LOGIN
# =========================================================

require_login()
import io
import re
import random
import zipfile
from copy import deepcopy
from pathlib import Path

import pandas as pd
import streamlit as st

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


# ============================================================
# PAGE CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="Change Skills + Add Projects",
    page_icon="🛠️",
    layout="wide",
)


# ============================================================
# PATHS
# ============================================================

THIS_FILE = Path(__file__).resolve()

PROJECT_ROOT = (
    THIS_FILE.parent.parent
    if THIS_FILE.parent.name.lower() == "pages"
    else THIS_FILE.parent
)

DATA_DIR = PROJECT_ROOT / "data"

CT_PROJECTS_FILE = (
    DATA_DIR / "CTProjects.xlsx"
)


# ============================================================
# SKILL OPTIONS
# ============================================================

SKILL_OPTIONS = [

    "C Programming",

    "C++ Programming",

    "Python Programming",

    "Python Full Stack",

    "Java Full Stack",

    "Gen AI",

    "Agentic AI",

    "Web Development",

    "MERN",

    "MEAN",

    "Cyber Security",

    "Cloud Computing",

    "DSA Using C",

    "DSA Using Python",

    "DSA Using Java",

    "Competitive Coding Using C",

    "Competitive Coding Using Python",

    "Competitive Coding Using Java",

    "Data Analytics using R",
]


# ============================================================
# COMMON TRAINING EXPERTISE
# ============================================================

COMMON_EXPERTISE = [

    "Mini Project Guidance & Development",

    "Major Project Mentoring & Guidance",

    "Problem Solving & Coding Practice",

    "Technical Assessments & Mock Interviews",

    "Student Mentoring & Project Guidance",
]


# ============================================================
# BASIC HELPERS
# ============================================================

def clean_text(value):

    if value is None:
        return ""

    try:

        if pd.isna(value):
            return ""

    except Exception:
        pass

    return str(value).strip()


def normalize_text(value):

    text = clean_text(value).lower()

    text = text.replace(
        "&",
        " and "
    )

    text = re.sub(
        r"[^a-z0-9+#.]+",
        " ",
        text
    )

    return re.sub(
        r"\s+",
        " ",
        text
    ).strip()


def normalized_heading(text):

    text = clean_text(text).lower()

    text = text.replace(
        ":",
        ""
    )

    return re.sub(
        r"\s+",
        " ",
        text
    ).strip()


# ============================================================
# SESSION STATE
# ============================================================

if "skills_projects_results" not in st.session_state:

    st.session_state.skills_projects_results = []


if "skills_projects_selected_skills" not in st.session_state:

    st.session_state.skills_projects_selected_skills = []


if "skills_projects_project_mode" not in st.session_state:

    st.session_state.skills_projects_project_mode = ""


if "skills_projects_project_count" not in st.session_state:

    st.session_state.skills_projects_project_count = 0


# ============================================================
# LOAD CT PROJECTS
# ============================================================

@st.cache_data(
    show_spinner=False
)
def load_project_pool(file_mtime):

    if not CT_PROJECTS_FILE.exists():

        raise FileNotFoundError(
            f"Not found: {CT_PROJECTS_FILE}"
        )


    try:

        df = pd.read_excel(
            CT_PROJECTS_FILE,
            sheet_name="Projects"
        )

    except ValueError:

        df = pd.read_excel(
            CT_PROJECTS_FILE
        )


    required = [

        "College / Client",

        "Location",

        "Domain / Subject Area",
    ]


    missing = [

        column

        for column in required

        if column not in df.columns
    ]


    if missing:

        raise ValueError(

            "CTProjects.xlsx is missing: "
            + ", ".join(missing)

        )


    df = df.copy()


    for column in required:

        df[column] = (
            df[column].apply(
                clean_text
            )
        )


    df = df[

        (df["College / Client"] != "")

        &

        (df["Location"] != "")

        &

        (df["Domain / Subject Area"] != "")

    ].copy()


    return (

        df.drop_duplicates(
            subset=required
        )

        .reset_index(
            drop=True
        )
    )


# ============================================================
# GET PARAGRAPHS
#
# Supports:
# Normal body paragraphs
# Paragraphs inside tables
# ============================================================

def get_paragraph_groups(doc):

    # Normal body
    yield doc.paragraphs


    # Tables
    for table in doc.tables:

        for row in table.rows:

            for cell in row.cells:

                yield cell.paragraphs


# ============================================================
# CLEAR PARAGRAPH CONTENT
#
# Keeps:
# Borders
# Spacing
# Alignment
# Indentation
# ============================================================

def clear_paragraph_keep_format(
        paragraph
):

    p = paragraph._p


    for child in list(p):

        if child.tag != qn(
            "w:pPr"
        ):

            p.remove(
                child
            )


# ============================================================
# REPLACE PARAGRAPH TEXT
#
# Preserves first run formatting.
# ============================================================

def replace_paragraph_text(
        paragraph,
        new_text
):

    run_format = None


    if paragraph.runs:

        first_run = (
            paragraph.runs[0]
        )


        if (
            first_run._r.rPr
            is not None
        ):

            run_format = deepcopy(
                first_run._r.rPr
            )


    clear_paragraph_keep_format(
        paragraph
    )


    new_run = paragraph.add_run(
        new_text
    )


    if run_format is not None:

        if new_run._r.rPr is not None:

            new_run._r.remove(
                new_run._r.rPr
            )


        new_run._r.insert(
            0,
            deepcopy(
                run_format
            )
        )


# ============================================================
# REMOVE PARAGRAPH
# ============================================================

def remove_paragraph(
        paragraph
):

    element = (
        paragraph._element
    )

    parent = (
        element.getparent()
    )


    if parent is not None:

        parent.remove(
            element
        )


# ============================================================
# CORE SKILLS TEXT
# ============================================================

def create_core_skills_text(
        skills
):

    return (

        "■ "

        + "    ■ ".join(
            skills
        )
    )


# ============================================================
# NATURAL LANGUAGE SKILL LIST
# ============================================================

def natural_skill_list(
        skills
):

    if not skills:

        return ""


    if len(skills) == 1:

        return skills[0]


    if len(skills) == 2:

        return (

            skills[0]

            + " and "

            + skills[1]
        )


    return (

        ", ".join(
            skills[:-1]
        )

        + ", and "

        + skills[-1]
    )


# ============================================================
# UPDATE CORE SKILLS
# ============================================================

def update_core_skills(
        doc,
        skills
):

    new_text = (
        create_core_skills_text(
            skills
        )
    )


    for paragraphs in (
        get_paragraph_groups(doc)
    ):


        for index, paragraph in enumerate(
            paragraphs
        ):


            if normalized_heading(
                paragraph.text
            ) != "core skills":

                continue


            # Find first non-empty paragraph
            # after Core Skills

            for next_index in range(

                index + 1,

                len(paragraphs)

            ):


                target = (
                    paragraphs[
                        next_index
                    ]
                )


                if target.text.strip():

                    replace_paragraph_text(

                        target,

                        new_text

                    )

                    return True


    return False


# ============================================================
# UPDATE PROFESSIONAL SUMMARY
#
# Only changes:
#
# Possesses strong expertise in ....
# ============================================================

def update_professional_summary(
        doc,
        skills
):

    new_sentence = (

        "Possesses strong expertise in "

        + natural_skill_list(
            skills
        )

        + "."
    )


    next_sections = {

        "professional highlights",

        "competitive coding profiles",

        "training expertise",

        "core competencies",

        "training projects",

        "completed projects",

        "ongoing projects",
    }


    pattern = (

        r"Possesses\s+strong\s+"

        r"expertise\s+in\s+.*?\."

    )


    for paragraphs in (
        get_paragraph_groups(doc)
    ):


        for index, paragraph in enumerate(
            paragraphs
        ):


            if normalized_heading(
                paragraph.text
            ) != "professional summary":

                continue


            for next_index in range(

                index + 1,

                len(paragraphs)

            ):


                summary_para = (

                    paragraphs[
                        next_index
                    ]
                )


                summary_text = (
                    summary_para.text.strip()
                )


                if not summary_text:

                    continue


                if normalized_heading(
                    summary_text
                ) in next_sections:

                    break


                if re.search(

                    pattern,

                    summary_para.text,

                    flags=re.IGNORECASE

                ):


                    updated_text = re.sub(

                        pattern,

                        new_sentence,

                        summary_para.text,

                        count=1,

                        flags=re.IGNORECASE

                    )


                    replace_paragraph_text(

                        summary_para,

                        updated_text

                    )


                    return True


    return False


# ============================================================
# SKILL → TRAINING EXPERTISE
# ============================================================

def skill_to_training_expertise(
        skill
):

    mapping = {

        "C Programming":
            "C Programming",

        "C++ Programming":
            "C++ Programming",

        "Python Programming":
            "Python Programming",

        "Python Full Stack":
            "Python Full Stack Training",

        "Java Full Stack":
            "Java Full Stack Training",

        "Gen AI":
            "Gen AI Training",

        "Agentic AI":
            "Agentic AI Training",

        "Web Development":
            "Web Development Training",

        "MERN":
            "MERN Stack Training",

        "MEAN":
            "MEAN Stack Training",

        "Cyber Security":
            "Cyber Security Training",

        "Cloud Computing":
            "Cloud Computing Training",

        "DSA Using C":
            "DSA Using C",

        "DSA Using Python":
            "DSA Using Python",

        "DSA Using Java":
            "DSA Using Java",

        "Competitive Coding Using C":
            "Competitive Coding Using C",

        "Competitive Coding Using Python":
            "Competitive Coding Using Python",

        "Competitive Coding Using Java":
            "Competitive Coding Using Java",

        "Data Analytics using R":
            "Data Analytics using R",
    }


    return mapping.get(
        skill,
        skill
    )


# ============================================================
# UPDATE TRAINING EXPERTISE
#
# IMPORTANT:
#
# All items are generated inside ONE paragraph.
#
# Line breaks are used instead of separate paragraphs.
#
# This prevents large spacing between expertise items.
# ============================================================

def update_training_expertise(
        doc,
        skills
):


    for paragraphs in (
        get_paragraph_groups(doc)
    ):


        heading_index = None

        next_section_index = None


        # ----------------------------------------------------
        # Find Training Expertise
        # and Core Competencies
        # ----------------------------------------------------

        for index, paragraph in enumerate(
            paragraphs
        ):


            heading = (
                normalized_heading(
                    paragraph.text
                )
            )


            if heading == "training expertise":

                heading_index = index

                continue


            if (

                heading_index
                is not None

                and

                heading
                == "core competencies"

            ):

                next_section_index = index

                break


        if (
            heading_index is None
            or
            next_section_index is None
        ):

            continue


        # ----------------------------------------------------
        # Existing expertise paragraphs
        # ----------------------------------------------------

        expertise_paragraphs = []


        for index in range(

            heading_index + 1,

            next_section_index

        ):


            if paragraphs[
                index
            ].text.strip():

                expertise_paragraphs.append(

                    paragraphs[
                        index
                    ]

                )


        if not expertise_paragraphs:

            continue


        # Use first old paragraph
        # to preserve formatting.

        target_para = (
            expertise_paragraphs[0]
        )


        expertise_items = []


        # ----------------------------------------------------
        # New skill-based expertise
        # ----------------------------------------------------

        for skill in skills:


            item = (

                skill_to_training_expertise(
                    skill
                )

            )


            if item not in expertise_items:

                expertise_items.append(
                    item
                )


        # ----------------------------------------------------
        # Common items
        # ----------------------------------------------------

        for item in COMMON_EXPERTISE:


            if item not in expertise_items:

                expertise_items.append(
                    item
                )


        # ----------------------------------------------------
        # ONE paragraph with line breaks
        # ----------------------------------------------------

        new_text = "\n".join(

            "● " + item

            for item in expertise_items

        )


        replace_paragraph_text(

            target_para,

            new_text

        )


        # Remove extra old paragraphs
        # if older profile used one paragraph/item.

        for extra_para in (
            expertise_paragraphs[1:]
        ):


            remove_paragraph(
                extra_para
            )


        return True


    return False


# ============================================================
# PROJECT MATCHING HELPERS
# ============================================================

DSA_TERMS = (

    "dsa",

    "data structure",

    "data structures",

    "algorithm",

    "algorithms",

    "competitive coding",

    "competitive programming",

    "problem solving",
)


def is_dsa(text):

    return any(

        term in text

        for term in DSA_TERMS

    )


def has_java(text):

    return bool(
        re.search(
            r"\bjava\b",
            text
        )
    )


def has_python(text):

    return bool(
        re.search(
            r"\bpython\b",
            text
        )
    )


def has_c_language(text):

    if (
        "c++" in text
        or
        "c plus plus" in text
        or
        "c#" in text
    ):

        return False


    return bool(

        re.search(
            r"\bc programming\b",
            text
        )

        or

        re.search(
            r"\bc language\b",
            text
        )

        or

        re.search(
            r"\bcore c\b",
            text
        )

        or

        re.search(
            r"\busing c\b",
            text
        )

        or

        re.search(
            r"\bc with dsa\b",
            text
        )

        or

        re.fullmatch(
            r"c",
            text
        )
    )


def has_cpp(text):

    return (

        "c++" in text

        or

        "c plus plus" in text

        or

        bool(
            re.search(
                r"\bcpp\b",
                text
            )
        )
    )


# ============================================================
# PROJECT MATCHING
# ============================================================

def project_matches_topic(
        subject,
        topic
):


    s = normalize_text(
        subject
    )

    t = normalize_text(
        topic
    )


    if not s or not t:

        return False


    # Strong direct match
    if t in s:

        return True


    # --------------------------------------------------------
    # C Programming
    # --------------------------------------------------------

    if t == "c programming":

        return has_c_language(
            s
        )


    # --------------------------------------------------------
    # C++
    # --------------------------------------------------------

    if t == "c++ programming":

        return has_cpp(
            s
        )


    # --------------------------------------------------------
    # Python Programming
    # --------------------------------------------------------

    if t == "python programming":

        return (

            has_python(s)

            and

            not is_dsa(s)

        )


    # --------------------------------------------------------
    # Python Full Stack
    # --------------------------------------------------------

    if t == "python full stack":

        return (

            (

                has_python(s)

                and

                any(

                    token in s

                    for token in (

                        "full stack",

                        "fullstack",

                        "django",

                        "flask",
                    )

                )
            )

            or

            "python full stack" in s
        )


    # --------------------------------------------------------
    # Java Full Stack
    # --------------------------------------------------------

    if t == "java full stack":

        return (

            (

                has_java(s)

                and

                any(

                    token in s

                    for token in (

                        "full stack",

                        "fullstack",

                        "spring",

                        "spring boot",
                    )

                )
            )

            or

            "java full stack" in s
        )


    # --------------------------------------------------------
    # Gen AI
    # --------------------------------------------------------

    if t == "gen ai":

        return any(

            token in s

            for token in (

                "gen ai",

                "genai",

                "generative ai",

                "prompt engineering",

                "llm",

                "large language model",
            )
        )


    # --------------------------------------------------------
    # Agentic AI
    # --------------------------------------------------------

    if t == "agentic ai":

        return any(

            token in s

            for token in (

                "agentic ai",

                "ai agent",

                "ai agents",

                "langgraph",

                "lang graph",

                "multi agent",
            )
        )


    # --------------------------------------------------------
    # Web Development
    # --------------------------------------------------------

    if t == "web development":

        return any(

            token in s

            for token in (

                "web development",

                "html",

                "css",

                "javascript",

                "react",

                "node",

                "express",

                "mern",

                "mean",

                "full stack",

                "fullstack",

                "django",

                "flask",
            )
        )


    # --------------------------------------------------------
    # MERN
    # --------------------------------------------------------

    if t == "mern":

        return any(

            token in s

            for token in (

                "mern",

                "mongo express react node",
            )
        )


    # --------------------------------------------------------
    # MEAN
    # --------------------------------------------------------

    if t == "mean":

        return any(

            token in s

            for token in (

                "mean",

                "mongo express angular node",
            )
        )


    # --------------------------------------------------------
    # Cyber Security
    # --------------------------------------------------------

    if t == "cyber security":

        return (

            "cyber security" in s

            or

            "cybersecurity" in s
        )


    # --------------------------------------------------------
    # Cloud Computing
    # --------------------------------------------------------

    if t == "cloud computing":

        return any(

            token in s

            for token in (

                "cloud",

                "aws",

                "azure",

                "gcp",
            )
        )


    # --------------------------------------------------------
    # DSA Using C
    # --------------------------------------------------------

    if t == "dsa using c":

        return (

            has_c_language(s)

            and

            is_dsa(s)
        )


    # --------------------------------------------------------
    # DSA Using Python
    # --------------------------------------------------------

    if t == "dsa using python":

        return (

            has_python(s)

            and

            is_dsa(s)
        )


    # --------------------------------------------------------
    # DSA Using Java
    # --------------------------------------------------------

    if t == "dsa using java":

        return (

            has_java(s)

            and

            is_dsa(s)
        )


    # --------------------------------------------------------
    # Competitive Coding Using C
    # --------------------------------------------------------

    if t == "competitive coding using c":

        return (

            has_c_language(s)

            and

            is_dsa(s)
        )


    # --------------------------------------------------------
    # Competitive Coding Using Python
    # --------------------------------------------------------

    if t == "competitive coding using python":

        return (

            has_python(s)

            and

            is_dsa(s)
        )


    # --------------------------------------------------------
    # Competitive Coding Using Java
    # --------------------------------------------------------

    if t == "competitive coding using java":

        return (

            has_java(s)

            and

            is_dsa(s)
        )


    # --------------------------------------------------------
    # Data Analytics using R
    # --------------------------------------------------------

    if t == "data analytics using r":

        return (

            (

                any(

                    token in s

                    for token in (

                        "data analytics",

                        "data analysis",

                        "analytics",
                    )
                )

                and

                bool(
                    re.search(
                        r"\br\b",
                        s
                    )
                )
            )

            or

            "data analytics using r" in s
        )


    # --------------------------------------------------------
    # Generic fallback
    # --------------------------------------------------------

    words = [

        word

        for word in t.split()

        if len(word) > 2
    ]


    return (

        bool(words)

        and

        all(
            word in s
            for word in words
        )
    )


# ============================================================
# MATCHING PROJECT ROWS
# ============================================================

def rows_for_topic(
        project_pool,
        topic,
        used_indexes=None
):


    used_indexes = (
        used_indexes
        or set()
    )


    matches = []


    for index, row in (
        project_pool.iterrows()
    ):


        if index in used_indexes:

            continue


        if project_matches_topic(

            row[
                "Domain / Subject Area"
            ],

            topic

        ):

            matches.append(
                index
            )


    return matches


# ============================================================
# GET REAL COLLEGE + LOCATION
# ============================================================

def pick_real_college_location(
        project_pool,
        used_pairs=None
):


    used_pairs = (
        used_pairs
        or set()
    )


    base = (

        project_pool[

            [
                "College / Client",
                "Location"
            ]

        ]

        .drop_duplicates()

        .to_dict(
            "records"
        )
    )


    random.shuffle(
        base
    )


    for row in base:


        pair = (

            clean_text(
                row[
                    "College / Client"
                ]
            ).lower(),

            clean_text(
                row[
                    "Location"
                ]
            ).lower(),
        )


        if pair not in used_pairs:

            return row


    return (

        random.choice(base)

        if base

        else None
    )


# ============================================================
# CREATE PROJECT IF MATCH DOES NOT EXIST
# ============================================================

def synthesize_project(
        project_pool,
        topic,
        used_pairs=None
):


    base = pick_real_college_location(

        project_pool,

        used_pairs

    )


    if not base:

        return None


    return {

        "College / Client":

            clean_text(
                base[
                    "College / Client"
                ]
            ),

        "Location":

            clean_text(
                base[
                    "Location"
                ]
            ),

        "Domain / Subject Area":

            clean_text(
                topic
            ),
    }


# ============================================================
# SELECT PROJECTS
# ============================================================

def select_projects(
        project_pool,
        count,
        mode,
        selected_skills
):


    count = int(
        count
    )


    used_indexes = set()

    used_pairs = set()

    rows = []

    synthesized_topics = []


    # ========================================================
    # RANDOM MODE
    # ========================================================

    if mode == "Random from all CTProjects":


        take = min(

            count,

            len(project_pool)

        )


        if take > 0:


            indexes = random.sample(

                list(
                    project_pool.index
                ),

                take

            )


            for index in indexes:


                row = (
                    project_pool.loc[
                        index
                    ]
                )


                rows.append(

                    {

                        "College / Client":

                            clean_text(
                                row[
                                    "College / Client"
                                ]
                            ),

                        "Location":

                            clean_text(
                                row[
                                    "Location"
                                ]
                            ),

                        "Domain / Subject Area":

                            clean_text(
                                row[
                                    "Domain / Subject Area"
                                ]
                            ),
                    }

                )


        topics_used = []


    # ========================================================
    # BASED ON SELECTED NEW SKILLS
    # ========================================================

    else:


        topics = list(
            selected_skills
        )


        if not topics:


            return (

                pd.DataFrame(

                    columns=[

                        "S.No",

                        "College / Client",

                        "Location",

                        "Domain / Subject Area",
                    ]
                ),

                [],

                [],
            )


        # Randomize topic order.
        # Then cycle through all selected topics.

        topic_cycle = list(
            topics
        )


        random.shuffle(
            topic_cycle
        )


        slot = 0

        safety = 0


        max_attempts = (

            count

            * max(

                10,

                len(topic_cycle)
                * 5

            )
        )


        while (

            len(rows) < count

            and

            safety < max_attempts

        ):


            topic = (

                topic_cycle[

                    slot
                    % len(topic_cycle)

                ]
            )


            slot += 1

            safety += 1


            # ------------------------------------------------
            # Find existing matching project
            # ------------------------------------------------

            matches = rows_for_topic(

                project_pool,

                topic,

                used_indexes

            )


            if matches:


                index = random.choice(
                    matches
                )


                source = (
                    project_pool.loc[
                        index
                    ]
                )


                rows.append(

                    {

                        "College / Client":

                            clean_text(
                                source[
                                    "College / Client"
                                ]
                            ),

                        "Location":

                            clean_text(
                                source[
                                    "Location"
                                ]
                            ),

                        "Domain / Subject Area":

                            clean_text(
                                source[
                                    "Domain / Subject Area"
                                ]
                            ),
                    }

                )


                used_indexes.add(
                    index
                )


                used_pairs.add(

                    (

                        clean_text(
                            source[
                                "College / Client"
                            ]
                        ).lower(),

                        clean_text(
                            source[
                                "Location"
                            ]
                        ).lower(),
                    )

                )


                continue


            # ------------------------------------------------
            # No matching project found.
            #
            # Use a real college and location from CTProjects
            # and selected skill as Domain.
            # ------------------------------------------------

            created = synthesize_project(

                project_pool,

                topic,

                used_pairs

            )


            if created:


                rows.append(
                    created
                )


                used_pairs.add(

                    (

                        created[
                            "College / Client"
                        ].lower(),

                        created[
                            "Location"
                        ].lower(),
                    )

                )


                if (
                    topic
                    not in synthesized_topics
                ):

                    synthesized_topics.append(
                        topic
                    )


        topics_used = topics


    # ========================================================
    # CREATE DATAFRAME
    # ========================================================

    result = pd.DataFrame(

        rows,

        columns=[

            "College / Client",

            "Location",

            "Domain / Subject Area",
        ],
    )


    if result.empty:


        return (

            pd.DataFrame(

                columns=[

                    "S.No",

                    "College / Client",

                    "Location",

                    "Domain / Subject Area",
                ]
            ),

            [],

            synthesized_topics,
        )


    result.insert(

        0,

        "S.No",

        range(
            1,
            len(result) + 1
        )
    )


    return (

        result,

        topics_used,

        synthesized_topics

    )


# ============================================================
# FIND WORD SECTION
# ============================================================

def find_section_paragraph(
        doc,
        needle
):


    needle = normalize_text(
        needle
    )


    for paragraph in doc.paragraphs:


        if needle in normalize_text(
            paragraph.text
        ):

            return paragraph


    return None


# ============================================================
# INSERT PARAGRAPH AFTER
# ============================================================

def insert_paragraph_after(
        doc,
        reference_paragraph,
        text
):


    paragraph = doc.add_paragraph()


    reference_paragraph._p.addnext(
        paragraph._p
    )


    run = paragraph.add_run(
        text
    )


    return (
        paragraph,
        run
    )


# ============================================================
# XML TEXT
# ============================================================

def xml_text(element):


    return "".join(

        node.text or ""

        for node in element.xpath(
            ".//w:t"
        )

    ).strip()


# ============================================================
# FIND LAST TABLE AFTER COMPLETED PROJECTS
# ============================================================

def last_table_immediately_after(
        heading_paragraph
):


    current = (
        heading_paragraph._p.getnext()
    )


    last_table = None


    while current is not None:


        tag = (
            current.tag
            .split("}")[-1]
        )


        if tag == "tbl":

            last_table = current


        elif tag == "p":

            if xml_text(
                current
            ):

                break


        else:

            break


        current = (
            current.getnext()
        )


    return last_table


# ============================================================
# TABLE BORDERS
# ============================================================

def set_table_borders(
        table,
        color="A6A6A6",
        size="6"
):


    tbl_pr = (
        table._tbl.tblPr
    )


    borders = tbl_pr.find(
        qn("w:tblBorders")
    )


    if borders is None:


        borders = OxmlElement(
            "w:tblBorders"
        )


        tbl_pr.append(
            borders
        )


    for edge in (

        "top",

        "left",

        "bottom",

        "right",

        "insideH",

        "insideV"

    ):


        node = borders.find(
            qn(
                f"w:{edge}"
            )
        )


        if node is None:


            node = OxmlElement(
                f"w:{edge}"
            )


            borders.append(
                node
            )


        node.set(
            qn("w:val"),
            "single"
        )


        node.set(
            qn("w:sz"),
            size
        )


        node.set(
            qn("w:space"),
            "0"
        )


        node.set(
            qn("w:color"),
            color
        )


# ============================================================
# TABLE CELL SHADING
# ============================================================

def set_cell_shading(
        cell,
        fill
):


    tc_pr = (
        cell._tc.get_or_add_tcPr()
    )


    node = tc_pr.find(
        qn("w:shd")
    )


    if node is None:


        node = OxmlElement(
            "w:shd"
        )


        tc_pr.append(
            node
        )


    node.set(
        qn("w:fill"),
        fill
    )


# ============================================================
# TABLE CELL MARGINS
# ============================================================

def set_cell_margins(
        cell,
        top=55,
        start=60,
        bottom=55,
        end=60
):


    tc_pr = (
        cell._tc.get_or_add_tcPr()
    )


    margins = tc_pr.find(
        qn("w:tcMar")
    )


    if margins is None:


        margins = OxmlElement(
            "w:tcMar"
        )


        tc_pr.append(
            margins
        )


    for name, value in {

        "top": top,

        "start": start,

        "bottom": bottom,

        "end": end,

    }.items():


        node = margins.find(
            qn(
                f"w:{name}"
            )
        )


        if node is None:


            node = OxmlElement(
                f"w:{name}"
            )


            margins.append(
                node
            )


        node.set(
            qn("w:w"),
            str(value)
        )


        node.set(
            qn("w:type"),
            "dxa"
        )


# ============================================================
# STYLE PROJECT TABLE
# ============================================================

def style_projects_table(
        table
):


    set_table_borders(
        table
    )


    table.autofit = True


    # --------------------------------------------------------
    # HEADER
    # --------------------------------------------------------

    for cell in (
        table.rows[0].cells
    ):


        set_cell_shading(
            cell,
            "E97132"
        )


        set_cell_margins(
            cell
        )


        cell.vertical_alignment = (
            WD_CELL_VERTICAL_ALIGNMENT.CENTER
        )


        for paragraph in (
            cell.paragraphs
        ):


            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )


            for run in (
                paragraph.runs
            ):


                run.bold = True


                run.font.size = Pt(
                    8.5
                )


                run.font.color.rgb = (
                    RGBColor(
                        255,
                        255,
                        255
                    )
                )


    # --------------------------------------------------------
    # BODY
    # --------------------------------------------------------

    for row in table.rows[1:]:


        for col_index, cell in enumerate(
            row.cells
        ):


            set_cell_margins(
                cell
            )


            cell.vertical_alignment = (
                WD_CELL_VERTICAL_ALIGNMENT.CENTER
            )


            for paragraph in (
                cell.paragraphs
            ):


                paragraph.alignment = (

                    WD_ALIGN_PARAGRAPH.CENTER

                    if col_index == 0

                    else

                    WD_ALIGN_PARAGRAPH.LEFT
                )


                for run in (
                    paragraph.runs
                ):


                    run.font.size = Pt(
                        8.5
                    )


# ============================================================
# ADD PROJECT TABLE
# ============================================================

def add_projects_table(
        doc,
        projects
):


    table = doc.add_table(
        rows=1,
        cols=4
    )


    headers = [

        "S.No",

        "College / Client",

        "Location",

        "Domain / Subject Area",
    ]


    for index, header in enumerate(
        headers
    ):


        table.rows[0].cells[
            index
        ].text = header


    for _, row in (
        projects.iterrows()
    ):


        cells = (
            table.add_row().cells
        )


        cells[0].text = str(
            row["S.No"]
        )


        cells[1].text = clean_text(
            row[
                "College / Client"
            ]
        )


        cells[2].text = clean_text(
            row[
                "Location"
            ]
        )


        cells[3].text = clean_text(
            row[
                "Domain / Subject Area"
            ]
        )


    style_projects_table(
        table
    )


    return table


# ============================================================
# INSERT PROJECTS INTO PROFILE
# ============================================================

def insert_projects_into_document(
        doc,
        projects
):


    # --------------------------------------------------------
    # Remove placeholder:
    # Project allocation not yet done
    # --------------------------------------------------------

    for paragraph in list(
        doc.paragraphs
    ):


        if (
            "project allocation not yet done"
            in normalize_text(
                paragraph.text
            )
        ):


            remove_paragraph(
                paragraph
            )


    # --------------------------------------------------------
    # TRAINING PROJECTS HEADING
    # --------------------------------------------------------

    training_heading = (
        find_section_paragraph(
            doc,
            "Training Projects"
        )
    )


    if training_heading is None:


        training_heading = (
            doc.add_paragraph()
        )


        run = training_heading.add_run(
            "Training Projects:"
        )


        run.bold = True


        run.font.size = Pt(
            12
        )


        run.font.color.rgb = (
            RGBColor(
                233,
                113,
                50
            )
        )


    # --------------------------------------------------------
    # COMPLETED PROJECTS HEADING
    # --------------------------------------------------------

    completed_heading = (
        find_section_paragraph(
            doc,
            "Completed Projects"
        )
    )


    if completed_heading is None:


        completed_heading, run = (
            insert_paragraph_after(

                doc,

                training_heading,

                "Completed Projects:"
            )
        )


        run.bold = True


        run.font.size = Pt(
            11
        )


        run.font.color.rgb = (
            RGBColor(
                233,
                113,
                50
            )
        )


        insert_after = (
            completed_heading._p
        )


    else:


        existing_table = (
            last_table_immediately_after(
                completed_heading
            )
        )


        insert_after = (

            existing_table

            if existing_table is not None

            else

            completed_heading._p
        )


    # --------------------------------------------------------
    # ADD TABLE
    # --------------------------------------------------------

    table = add_projects_table(

        doc,

        projects

    )


    insert_after.addnext(
        table._tbl
    )


    return True


# ============================================================
# PROCESS ONE WORD FILE
# ============================================================

def process_word_file(

        uploaded_file,

        selected_skills,

        project_count,

        project_mode,

        project_pool

):


    doc = Document(

        io.BytesIO(
            uploaded_file.getvalue()
        )

    )


    # ========================================================
    # CHANGE SKILLS
    # ========================================================

    section_results = {


        "Core Skills":

            update_core_skills(
                doc,
                selected_skills
            ),


        "Professional Summary":

            update_professional_summary(
                doc,
                selected_skills
            ),


        "Training Expertise":

            update_training_expertise(
                doc,
                selected_skills
            ),
    }


    # ========================================================
    # SELECT PROJECTS
    # ========================================================

    projects, topics_used, synthesized_topics = (

        select_projects(

            project_pool=project_pool,

            count=project_count,

            mode=project_mode,

            selected_skills=selected_skills,

        )

    )


    # ========================================================
    # ADD PROJECTS
    # ========================================================

    projects_added = False


    if not projects.empty:


        projects_added = (
            insert_projects_into_document(

                doc,

                projects

            )
        )


    # ========================================================
    # SAVE WORD FILE
    # ========================================================

    stream = io.BytesIO()


    doc.save(
        stream
    )


    stem = Path(
        uploaded_file.name
    ).stem


    return {


        "success":

            all(
                section_results.values()
            )

            and

            projects_added,


        "filename":

            uploaded_file.name,


        "output_name":

            (
                f"{stem}"
                f"_Skills_Projects_Updated.docx"
            ),


        "output_bytes":

            stream.getvalue(),


        "section_results":

            section_results,


        "projects_added":

            projects_added,


        "projects":

            projects,


        "topics_used":

            topics_used,


        "synthesized_topics":

            synthesized_topics,


        "project_mode":

            project_mode,
    }


# ============================================================
# PAGE TITLE
# ============================================================

st.title(
    "🛠️ Change Skills + Add Projects"
)


st.caption(

    "Upload trainer Word profiles, select the new skills, "
    "and add completed projects based on those selected skills."

)


# ============================================================
# CHECK CTPROJECTS FILE
# ============================================================

if CT_PROJECTS_FILE.exists():


    try:


        pool_count = len(

            load_project_pool(

                CT_PROJECTS_FILE.stat().st_mtime

            )

        )


        st.success(

            f"✅ CTProjects.xlsx found — "
            f"{pool_count} project records"

        )


    except Exception as exc:


        st.error(

            f"❌ Unable to read CTProjects.xlsx: {exc}"

        )


else:


    st.error(

        "❌ data/CTProjects.xlsx not found"

    )


# ============================================================
# 1. UPLOAD FILES
# ============================================================

st.markdown(
    "### 1. Upload Trainer Word Files"
)


uploaded_files = st.file_uploader(

    "Upload one or more .docx trainer profiles",

    type=["docx"],

    accept_multiple_files=True,

)


if uploaded_files:


    st.success(

        f"✅ {len(uploaded_files)} file(s) selected"

    )


    with st.expander(
        "📄 View Selected Files"
    ):


        for index, uploaded_file in enumerate(

            uploaded_files,

            start=1

        ):


            st.write(

                f"{index}. {uploaded_file.name}"

            )


# ============================================================
# 2. SELECT NEW SKILLS
# ============================================================

st.markdown(
    "### 2. Select New Skills"
)


st.caption(

    "The selected skills will replace the existing Core Skills "
    "and will also update Professional Summary and Training Expertise."

)


selected_skills = []


# ============================================================
# CHECKBOXES IN 3 COLUMNS
# ============================================================

skill_columns = st.columns(
    3
)


for index, skill in enumerate(
    SKILL_OPTIONS
):


    with skill_columns[
        index % 3
    ]:


        checked = st.checkbox(

            skill,

            key=f"combined_skill_{index}"

        )


        if checked:


            selected_skills.append(
                skill
            )


# ============================================================
# SHOW SELECTED SKILLS
# ============================================================

if selected_skills:


    st.info(

        "**Selected Skills:**  "

        + "  ■ ".join(
            selected_skills
        )

    )


    st.caption(

        f"{len(selected_skills)} skill(s) selected"

    )


else:


    st.warning(

        "Select at least one skill."

    )


# ============================================================
# 3. PROJECT SETTINGS
# ============================================================

st.markdown(
    "### 3. Project Settings"
)


project_col1, project_col2 = (
    st.columns(2)
)


with project_col1:


    number_of_projects = st.number_input(

        "Number of projects to add to EACH trainer",

        min_value=1,

        max_value=30,

        value=5,

        step=1,

    )


with project_col2:


    project_mode = st.radio(

        "Project Selection",

        options=[

            "Based on Selected Skills",

            "Random from all CTProjects",

        ],

        index=0,

    )


# ============================================================
# PROJECT MODE MESSAGE
# ============================================================

if project_mode == "Based on Selected Skills":


    st.info(

        "Projects will be selected using the NEW skills checked above. "
        "If a matching project is unavailable, the app uses a real "
        "College / Client and Location from CTProjects.xlsx and uses "
        "the selected skill as the Domain / Subject Area."

    )


else:


    st.info(

        "Projects will be selected randomly from the complete "
        "CTProjects.xlsx project pool."

    )


# ============================================================
# 4. PROCESS PROFILES
# ============================================================

st.markdown(
    "### 4. Process Profiles"
)


process_disabled = (

    not uploaded_files

    or

    not selected_skills

    or

    not CT_PROJECTS_FILE.exists()

)


if st.button(

    "🚀 Change Skills + Add Projects",

    type="primary",

    use_container_width=True,

    disabled=process_disabled,

):


    try:


        project_pool = (
            load_project_pool(

                CT_PROJECTS_FILE.stat().st_mtime

            )
        )


        results = []


        progress = st.progress(
            0
        )


        status = st.empty()


        total_files = len(
            uploaded_files
        )


        # ====================================================
        # PROCESS EACH FILE
        # ====================================================

        for position, uploaded_file in enumerate(

            uploaded_files,

            start=1

        ):


            status.write(

                f"Processing {position} of "
                f"{total_files}: "
                f"{uploaded_file.name}"

            )


            result = process_word_file(

                uploaded_file=uploaded_file,

                selected_skills=selected_skills,

                project_count=number_of_projects,

                project_mode=project_mode,

                project_pool=project_pool,

            )


            results.append(
                result
            )


            progress.progress(

                position
                / total_files

            )


        status.empty()


        # ====================================================
        # SAVE RESULTS TO SESSION STATE
        # ====================================================

        st.session_state.skills_projects_results = (
            results
        )


        st.session_state.skills_projects_selected_skills = (

            selected_skills.copy()

        )


        st.session_state.skills_projects_project_mode = (

            project_mode

        )


        st.session_state.skills_projects_project_count = (

            int(
                number_of_projects
            )

        )


        # ====================================================
        # SUCCESS COUNT
        # ====================================================

        success_count = sum(

            1

            for result in results

            if result.get(
                "success"
            )

        )


        st.success(

            f"Completed: {success_count} of "
            f"{len(results)} file(s) fully updated."

        )


    except Exception as exc:


        st.exception(
            exc
        )


# ============================================================
# 5. RESULTS / DOWNLOADS
#
# OUTSIDE BUTTON BLOCK:
#
# Downloading one file will NOT remove remaining files.
# ============================================================

results = (
    st.session_state.skills_projects_results
)


if results:


    st.markdown(
        "---"
    )


    st.markdown(
        "### 5. Updated Trainer Profiles"
    )


    # ========================================================
    # SHOW SKILLS USED
    # ========================================================

    if (
        st.session_state
        .skills_projects_selected_skills
    ):


        st.info(

            "**Skills Applied:**  "

            + "  ■ ".join(

                st.session_state
                .skills_projects_selected_skills

            )

        )


    # ========================================================
    # SHOW PROJECT SETTINGS USED
    # ========================================================

    st.write(

        "**Project Selection:**",

        st.session_state
        .skills_projects_project_mode,

    )


    st.write(

        "**Projects per Trainer:**",

        st.session_state
        .skills_projects_project_count,

    )


    # ========================================================
    # STATUS TABLE
    # ========================================================

    status_rows = []


    for result in results:


        sections = result.get(

            "section_results",

            {}

        )


        status_rows.append(

            {

                "File":

                    result.get(
                        "filename",
                        ""
                    ),

                "Core Skills":

                    "✅"

                    if sections.get(
                        "Core Skills"
                    )

                    else "❌",

                "Professional Summary":

                    "✅"

                    if sections.get(
                        "Professional Summary"
                    )

                    else "❌",

                "Training Expertise":

                    "✅"

                    if sections.get(
                        "Training Expertise"
                    )

                    else "❌",

                "Projects":

                    "✅"

                    if result.get(
                        "projects_added"
                    )

                    else "❌",

                "Status":

                    "Updated"

                    if result.get(
                        "success"
                    )

                    else

                    "Updated with warning",

            }

        )


    st.dataframe(

        status_rows,

        hide_index=True,

        use_container_width=True,

    )


    # ========================================================
    # DOWNLOAD ALL ZIP
    # ========================================================

    downloadable = [

        result

        for result in results

        if result.get(
            "output_bytes"
        )

    ]


    if downloadable:


        zip_stream = io.BytesIO()


        with zipfile.ZipFile(

            zip_stream,

            "w",

            zipfile.ZIP_DEFLATED,

        ) as archive:


            for result in downloadable:


                archive.writestr(

                    result[
                        "output_name"
                    ],

                    result[
                        "output_bytes"
                    ],

                )


        st.download_button(

            "📦 Download All Updated Profiles as ZIP",

            data=zip_stream.getvalue(),

            file_name=(
                "Trainer_Profiles_"
                "Skills_Projects_Updated.zip"
            ),

            mime="application/zip",

            use_container_width=True,

            type="primary",

            key="combined_download_all",

            on_click="ignore",

        )


    # ========================================================
    # INDIVIDUAL FILES
    # ========================================================

    st.markdown(
        "#### Individual Files"
    )


    for index, result in enumerate(

        results,

        start=1

    ):


        title = (

            result.get(
                "filename"
            )

            or

            f"File {index}"

        )


        with st.expander(

            f"{index}. {title}",

            expanded=(
                len(results) <= 5
            ),

        ):


            sections = result.get(

                "section_results",

                {}

            )


            col1, col2, col3, col4 = (

                st.columns(4)

            )


            with col1:


                st.write(

                    "**Core Skills:**",

                    "✅"

                    if sections.get(
                        "Core Skills"
                    )

                    else "❌",

                )


            with col2:


                st.write(

                    "**Summary:**",

                    "✅"

                    if sections.get(
                        "Professional Summary"
                    )

                    else "❌",

                )


            with col3:


                st.write(

                    "**Training Expertise:**",

                    "✅"

                    if sections.get(
                        "Training Expertise"
                    )

                    else "❌",

                )


            with col4:


                st.write(

                    "**Projects:**",

                    "✅"

                    if result.get(
                        "projects_added"
                    )

                    else "❌",

                )


            # =================================================
            # TECHNOLOGIES USED
            # =================================================

            if result.get(
                "topics_used"
            ):


                st.write(

                    "**Technologies used for project selection:**",

                    ", ".join(

                        result[
                            "topics_used"
                        ]

                    ),

                )


            # =================================================
            # SYNTHESIZED PROJECT WARNING
            # =================================================

            if result.get(
                "synthesized_topics"
            ):


                st.warning(

                    "No matching / sufficient project was found for: "

                    + ", ".join(

                        result[
                            "synthesized_topics"
                        ]

                    )

                    + ". A real College / Client and Location from "
                    "CTProjects.xlsx was used with the selected "
                    "technology as Domain / Subject Area."

                )


            # =================================================
            # PROJECT TABLE PREVIEW
            # =================================================

            projects = result.get(
                "projects"
            )


            if (

                isinstance(
                    projects,
                    pd.DataFrame
                )

                and

                not projects.empty

            ):


                st.dataframe(

                    projects[

                        [

                            "S.No",

                            "College / Client",

                            "Location",

                            "Domain / Subject Area",
                        ]

                    ],

                    hide_index=True,

                    use_container_width=True,

                )


            # =================================================
            # INDIVIDUAL DOWNLOAD
            # =================================================

            if result.get(
                "output_bytes"
            ):


                st.download_button(

                    f"⬇️ Download "
                    f"{result['output_name']}",

                    data=result[
                        "output_bytes"
                    ],

                    file_name=result[
                        "output_name"
                    ],

                    mime=(

                        "application/vnd."
                        "openxmlformats-officedocument."
                        "wordprocessingml.document"

                    ),

                    key=(

                        f"combined_download_"
                        f"{index}_"
                        f"{result['output_name']}"

                    ),

                    use_container_width=True,

                    on_click="ignore",

                )


    # ========================================================
    # CLEAR GENERATED FILES
    # ========================================================

    st.markdown(
        "---"
    )


    if st.button(

        "🗑️ Clear Generated Files",

        use_container_width=True,

        key="clear_combined_results",

    ):


        st.session_state.skills_projects_results = []


        st.session_state.skills_projects_selected_skills = []


        st.session_state.skills_projects_project_mode = ""


        st.session_state.skills_projects_project_count = 0


        st.rerun()