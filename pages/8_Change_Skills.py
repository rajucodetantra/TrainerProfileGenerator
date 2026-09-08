import streamlit as st

from login import require_login


# =========================================================
# PAGE CONFIG
# =========================================================

st.set_page_config(
    page_title="Page Name",
    page_icon="📄",
    layout="wide"
)


# =========================================================
# LOGIN
# =========================================================

require_login()
import streamlit as st

from docx import Document
from io import BytesIO
from pathlib import Path
from copy import deepcopy
from docx.oxml.ns import qn

import zipfile
import re


# =========================================================
# PAGE CONFIGURATION
# =========================================================

st.set_page_config(
    page_title="Change Skills",
    page_icon="🛠️",
    layout="wide"
)

st.title("🛠️ Change Trainer Skills")

st.write(
    "Upload one or more Trainer Profile DOCX files, "
    "select the required skills and update all profiles."
)


# =========================================================
# SESSION STATE
# =========================================================

if "processed_files" not in st.session_state:
    st.session_state.processed_files = []

if "status_rows" not in st.session_state:
    st.session_state.status_rows = []

if "skills_processed" not in st.session_state:
    st.session_state.skills_processed = []


# =========================================================
# AVAILABLE SKILLS
# =========================================================

SKILL_OPTIONS = [

    "C Programming",
    "C++ Programming",

    "Python Programming",
    "Python Full Stack",

    "Java Full Stack",

    "Gen AI",
    "Agentic AI",

    "Web Development",

    "MERN",
    "MEAN",

    "Cyber Security",
    "Cloud Computing",

    "DSA Using C",
    "DSA Using Python",
    "DSA Using Java",

    "Competitive Coding Using C",
    "Competitive Coding Using Python",
    "Competitive Coding Using Java",

    "Data Analytics using R"
]


# =========================================================
# COMMON TRAINING EXPERTISE
#
# These items are retained for every trainer.
# =========================================================

COMMON_EXPERTISE = [

    "Mini Project Guidance & Development",

    "Major Project Mentoring & Guidance",

    "Problem Solving & Coding Practice",

    "Technical Assessments & Mock Interviews",

    "Student Mentoring & Project Guidance"
]


# =========================================================
# GET ALL PARAGRAPH GROUPS
#
# Supports:
# 1. Normal body paragraphs
# 2. Paragraphs inside tables
# =========================================================

def get_paragraph_groups(doc):

    # Main document
    yield doc.paragraphs

    # Tables
    for table in doc.tables:

        for row in table.rows:

            for cell in row.cells:

                yield cell.paragraphs


# =========================================================
# NORMALIZE HEADING
# =========================================================

def normalized_heading(text):

    text = text.strip().lower()

    text = text.replace(":", "")

    text = re.sub(
        r"\s+",
        " ",
        text
    )

    return text


# =========================================================
# CLEAR PARAGRAPH CONTENT
#
# Keeps paragraph properties such as:
# - Borders
# - Spacing
# - Alignment
# - Indentation
# =========================================================

def clear_paragraph_keep_format(paragraph):

    p = paragraph._p

    for child in list(p):

        if child.tag != qn("w:pPr"):

            p.remove(child)


# =========================================================
# REPLACE PARAGRAPH TEXT
#
# Preserves:
# - Font
# - Font size
# - Bold
# - Italic
# - Paragraph spacing
# - Borders
# =========================================================

def replace_paragraph_text(
        paragraph,
        new_text
):

    run_format = None

    # Save first run formatting
    if paragraph.runs:

        first_run = paragraph.runs[0]

        if first_run._r.rPr is not None:

            run_format = deepcopy(
                first_run._r.rPr
            )

    # Remove existing content
    clear_paragraph_keep_format(
        paragraph
    )

    # Add updated text
    new_run = paragraph.add_run(
        new_text
    )

    # Restore character formatting
    if run_format is not None:

        if new_run._r.rPr is not None:

            new_run._r.remove(
                new_run._r.rPr
            )

        new_run._r.insert(
            0,
            deepcopy(run_format)
        )


# =========================================================
# DELETE PARAGRAPH
# =========================================================

def delete_paragraph(paragraph):

    element = paragraph._element

    parent = element.getparent()

    if parent is not None:

        parent.remove(element)


# =========================================================
# CREATE CORE SKILLS TEXT
#
# Example:
#
# ■ Python Programming    ■ Python Full Stack
# =========================================================

def create_core_skills_text(skills):

    return (
        "■ "
        + "    ■ ".join(skills)
    )


# =========================================================
# CREATE NATURAL ENGLISH SKILL LIST
#
# 1 Skill:
# Python Programming
#
# 2 Skills:
# Python Programming and Java Full Stack
#
# 3+ Skills:
# Python Programming, Java Full Stack, and Gen AI
# =========================================================

def natural_skill_list(skills):

    if not skills:

        return ""

    if len(skills) == 1:

        return skills[0]

    if len(skills) == 2:

        return (
            skills[0]
            + " and "
            + skills[1]
        )

    return (
        ", ".join(
            skills[:-1]
        )
        + ", and "
        + skills[-1]
    )


# =========================================================
# UPDATE CORE SKILLS
# =========================================================

def update_core_skills(
        doc,
        skills
):

    new_text = create_core_skills_text(
        skills
    )

    for paragraphs in get_paragraph_groups(doc):

        for i, paragraph in enumerate(paragraphs):

            heading = normalized_heading(
                paragraph.text
            )

            if heading == "core skills":

                # Find first non-empty paragraph
                # after Core Skills heading
                for j in range(
                    i + 1,
                    len(paragraphs)
                ):

                    target_para = paragraphs[j]

                    if target_para.text.strip():

                        replace_paragraph_text(
                            target_para,
                            new_text
                        )

                        return True

    return False


# =========================================================
# UPDATE PROFESSIONAL SUMMARY
#
# Only this sentence changes:
#
# Possesses strong expertise in .....
#
# Everything before and after this sentence stays unchanged.
# =========================================================

def update_professional_summary(
        doc,
        skills
):

    skill_text = natural_skill_list(
        skills
    )

    new_sentence = (
        "Possesses strong expertise in "
        + skill_text
        + "."
    )

    next_sections = {

        "professional highlights",
        "competitive coding profiles",
        "training expertise",
        "core competencies",
        "training projects",
        "completed projects",
        "ongoing projects"
    }

    for paragraphs in get_paragraph_groups(doc):

        for i, paragraph in enumerate(paragraphs):

            heading = normalized_heading(
                paragraph.text
            )

            if heading != "professional summary":

                continue

            # Search paragraphs after heading
            for j in range(
                i + 1,
                len(paragraphs)
            ):

                summary_para = paragraphs[j]

                summary_text = (
                    summary_para.text.strip()
                )

                if not summary_text:

                    continue

                possible_heading = (
                    normalized_heading(
                        summary_text
                    )
                )

                # Stop if next section begins
                if possible_heading in next_sections:

                    break

                # Find:
                # Possesses strong expertise in .....
                pattern = (
                    r"Possesses\s+strong\s+"
                    r"expertise\s+in\s+.*?\."
                )

                if re.search(
                    pattern,
                    summary_para.text,
                    flags=re.IGNORECASE
                ):

                    updated_text = re.sub(
                        pattern,
                        new_sentence,
                        summary_para.text,
                        count=1,
                        flags=re.IGNORECASE
                    )

                    replace_paragraph_text(
                        summary_para,
                        updated_text
                    )

                    return True

    return False


# =========================================================
# CONVERT CORE SKILL TO TRAINING EXPERTISE
# =========================================================

def skill_to_training_expertise(skill):

    # -----------------------------------------------------
    # Programming
    # -----------------------------------------------------

    if skill == "C Programming":

        return "C Programming"

    if skill == "C++ Programming":

        return "C++ Programming"

    if skill == "Python Programming":

        return "Python Programming"


    # -----------------------------------------------------
    # FULL STACK
    # -----------------------------------------------------

    if skill == "Python Full Stack":

        return "Python Full Stack Training"

    if skill == "Java Full Stack":

        return "Java Full Stack Training"


    # -----------------------------------------------------
    # AI
    # -----------------------------------------------------

    if skill == "Gen AI":

        return "Gen AI Training"

    if skill == "Agentic AI":

        return "Agentic AI Training"


    # -----------------------------------------------------
    # WEB
    # -----------------------------------------------------

    if skill == "Web Development":

        return "Web Development Training"


    # -----------------------------------------------------
    # MERN / MEAN
    # -----------------------------------------------------

    if skill == "MERN":

        return "MERN Stack Training"

    if skill == "MEAN":

        return "MEAN Stack Training"


    # -----------------------------------------------------
    # CYBER SECURITY
    # -----------------------------------------------------

    if skill == "Cyber Security":

        return "Cyber Security Training"


    # -----------------------------------------------------
    # CLOUD
    # -----------------------------------------------------

    if skill == "Cloud Computing":

        return "Cloud Computing Training"


    # -----------------------------------------------------
    # DSA
    # -----------------------------------------------------

    if skill == "DSA Using C":

        return "DSA Using C"

    if skill == "DSA Using Python":

        return "DSA Using Python"

    if skill == "DSA Using Java":

        return "DSA Using Java"


    # -----------------------------------------------------
    # COMPETITIVE CODING
    # -----------------------------------------------------

    if skill == "Competitive Coding Using C":

        return "Competitive Coding Using C"

    if skill == "Competitive Coding Using Python":

        return "Competitive Coding Using Python"

    if skill == "Competitive Coding Using Java":

        return "Competitive Coding Using Java"


    # -----------------------------------------------------
    # DATA ANALYTICS
    # -----------------------------------------------------

    if skill == "Data Analytics using R":

        return "Data Analytics using R"


    # -----------------------------------------------------
    # DEFAULT
    # -----------------------------------------------------

    return skill


# =========================================================
# UPDATE TRAINING EXPERTISE
#
# IMPORTANT:
#
# All expertise items are written inside ONE paragraph
# using line breaks.
#
# This keeps spacing compact like the existing template.
# =========================================================

def update_training_expertise(
        doc,
        skills
):

    for paragraphs in get_paragraph_groups(doc):

        heading_index = None
        next_section_index = None


        # -------------------------------------------------
        # FIND TRAINING EXPERTISE
        # -------------------------------------------------

        for i, paragraph in enumerate(paragraphs):

            heading = normalized_heading(
                paragraph.text
            )

            if heading == "training expertise":

                heading_index = i

                continue


            # Core Competencies marks end of
            # Training Expertise
            if (
                heading_index is not None
                and heading == "core competencies"
            ):

                next_section_index = i

                break


        if heading_index is None:

            continue


        if next_section_index is None:

            continue


        # -------------------------------------------------
        # GET OLD TRAINING EXPERTISE CONTENT
        # -------------------------------------------------

        expertise_paragraphs = []

        for i in range(
            heading_index + 1,
            next_section_index
        ):

            if paragraphs[i].text.strip():

                expertise_paragraphs.append(
                    paragraphs[i]
                )


        if not expertise_paragraphs:

            continue


        # -------------------------------------------------
        # USE FIRST OLD PARAGRAPH
        #
        # This preserves the original paragraph format.
        # -------------------------------------------------

        target_para = expertise_paragraphs[0]


        # -------------------------------------------------
        # BUILD NEW EXPERTISE LIST
        # -------------------------------------------------

        expertise_items = []


        # Skill-related expertise
        for skill in skills:

            item = (
                skill_to_training_expertise(
                    skill
                )
            )

            if item not in expertise_items:

                expertise_items.append(
                    item
                )


        # Common expertise
        for item in COMMON_EXPERTISE:

            if item not in expertise_items:

                expertise_items.append(
                    item
                )


        # -------------------------------------------------
        # CREATE ONE PARAGRAPH WITH LINE BREAKS
        #
        # Example:
        #
        # ● C Programming
        # ● Python Programming
        # ● Python Full Stack Training
        # ● DSA Using Python
        # -------------------------------------------------

        new_text = "\n".join(

            "● " + item

            for item in expertise_items
        )


        # -------------------------------------------------
        # REPLACE FIRST OLD PARAGRAPH
        # -------------------------------------------------

        replace_paragraph_text(
            target_para,
            new_text
        )


        # -------------------------------------------------
        # REMOVE EXTRA OLD EXPERTISE PARAGRAPHS
        #
        # Handles old profiles where each expertise
        # item was stored separately.
        # -------------------------------------------------

        for extra_para in (
            expertise_paragraphs[1:]
        ):

            delete_paragraph(
                extra_para
            )


        return True


    return False


# =========================================================
# PROCESS ONE DOCUMENT
# =========================================================

def process_document(
        file_bytes,
        skills
):

    doc = Document(
        BytesIO(file_bytes)
    )


    results = {

        "Core Skills":
            False,

        "Professional Summary":
            False,

        "Training Expertise":
            False
    }


    # -----------------------------------------------------
    # UPDATE CORE SKILLS
    # -----------------------------------------------------

    results["Core Skills"] = (
        update_core_skills(
            doc,
            skills
        )
    )


    # -----------------------------------------------------
    # UPDATE PROFESSIONAL SUMMARY
    # -----------------------------------------------------

    results["Professional Summary"] = (
        update_professional_summary(
            doc,
            skills
        )
    )


    # -----------------------------------------------------
    # UPDATE TRAINING EXPERTISE
    # -----------------------------------------------------

    results["Training Expertise"] = (
        update_training_expertise(
            doc,
            skills
        )
    )


    # -----------------------------------------------------
    # SAVE UPDATED DOCUMENT
    # -----------------------------------------------------

    output = BytesIO()

    doc.save(
        output
    )

    output.seek(0)


    return (
        output.getvalue(),
        results
    )


# =========================================================
# MULTIPLE FILE UPLOAD
# =========================================================

uploaded_files = st.file_uploader(
    "Upload Trainer Profile Files",
    type=["docx"],
    accept_multiple_files=True
)


# =========================================================
# SHOW SELECTED FILES
# =========================================================

if uploaded_files:

    st.success(
        f"✅ {len(uploaded_files)} file(s) selected"
    )


    with st.expander(
        "📄 View Selected Files"
    ):

        for number, file in enumerate(
            uploaded_files,
            start=1
        ):

            st.write(
                f"{number}. {file.name}"
            )


# =========================================================
# SKILL CHECKBOXES
# =========================================================

st.subheader(
    "Select New Skills"
)

st.caption(
    "Select all skills that should replace the existing "
    "skills in the uploaded trainer profiles."
)


selected_skills = []


# =========================================================
# DISPLAY CHECKBOXES IN 3 COLUMNS
# =========================================================

col1, col2, col3 = st.columns(3)

columns = [
    col1,
    col2,
    col3
]


for index, skill in enumerate(
    SKILL_OPTIONS
):

    current_column = (
        columns[
            index % 3
        ]
    )


    with current_column:

        selected = st.checkbox(
            skill,
            key=f"trainer_skill_{index}"
        )


        if selected:

            selected_skills.append(
                skill
            )


# =========================================================
# SELECTED SKILLS PREVIEW
# =========================================================

if selected_skills:

    st.divider()

    st.write(
        "**Selected Skills:**"
    )

    st.info(
        "■ "
        + "    ■ ".join(
            selected_skills
        )
    )

    st.caption(
        f"{len(selected_skills)} skill(s) selected"
    )


# =========================================================
# PROCESS ALL FILES
# =========================================================

if st.button(
    "🔄 Change Skills in All Profiles",
    type="primary",
    use_container_width=True
):


    # -----------------------------------------------------
    # FILE VALIDATION
    # -----------------------------------------------------

    if not uploaded_files:

        st.warning(
            "Please upload at least one DOCX file."
        )


    # -----------------------------------------------------
    # SKILL VALIDATION
    # -----------------------------------------------------

    elif not selected_skills:

        st.warning(
            "Please select at least one skill."
        )


    else:

        skills = selected_skills.copy()


        # -------------------------------------------------
        # CLEAR PREVIOUS GENERATED RESULTS
        # -------------------------------------------------

        st.session_state.processed_files = []

        st.session_state.status_rows = []

        st.session_state.skills_processed = (
            skills.copy()
        )


        # -------------------------------------------------
        # PROGRESS
        # -------------------------------------------------

        total_files = len(
            uploaded_files
        )

        progress_bar = st.progress(
            0
        )

        status_text = st.empty()


        # =================================================
        # PROCESS EACH FILE
        # =================================================

        for index, uploaded_file in enumerate(
            uploaded_files,
            start=1
        ):


            status_text.write(

                f"Processing {index} of "
                f"{total_files}: "
                f"{uploaded_file.name}"

            )


            try:

                # -----------------------------------------
                # PROCESS DOCUMENT
                # -----------------------------------------

                updated_bytes, results = (
                    process_document(
                        uploaded_file.getvalue(),
                        skills
                    )
                )


                # -----------------------------------------
                # OUTPUT FILE NAME
                # -----------------------------------------

                original_name = Path(
                    uploaded_file.name
                ).stem


                output_name = (

                    f"{original_name}"
                    f"_Skills_Updated.docx"

                )


                # -----------------------------------------
                # SAVE FILE IN SESSION STATE
                # -----------------------------------------

                st.session_state.processed_files.append(

                    {

                        "name":
                            output_name,

                        "data":
                            updated_bytes

                    }

                )


                # -----------------------------------------
                # CHECK SECTION RESULTS
                # -----------------------------------------

                core_status = (

                    "✅"

                    if results[
                        "Core Skills"
                    ]

                    else "❌"

                )


                summary_status = (

                    "✅"

                    if results[
                        "Professional Summary"
                    ]

                    else "❌"

                )


                expertise_status = (

                    "✅"

                    if results[
                        "Training Expertise"
                    ]

                    else "❌"

                )


                # -----------------------------------------
                # OVERALL STATUS
                # -----------------------------------------

                if all(
                    results.values()
                ):

                    overall_status = (
                        "Updated"
                    )

                else:

                    overall_status = (
                        "Updated with warning"
                    )


                # -----------------------------------------
                # STORE STATUS
                # -----------------------------------------

                st.session_state.status_rows.append(

                    {

                        "File":
                            uploaded_file.name,

                        "Core Skills":
                            core_status,

                        "Professional Summary":
                            summary_status,

                        "Training Expertise":
                            expertise_status,

                        "Status":
                            overall_status

                    }

                )


            except Exception as e:


                st.session_state.status_rows.append(

                    {

                        "File":
                            uploaded_file.name,

                        "Core Skills":
                            "❌",

                        "Professional Summary":
                            "❌",

                        "Training Expertise":
                            "❌",

                        "Status":
                            f"Error: {str(e)}"

                    }

                )


            # ---------------------------------------------
            # UPDATE PROGRESS
            # ---------------------------------------------

            progress_bar.progress(
                index / total_files
            )


        # -------------------------------------------------
        # REMOVE PROCESSING MESSAGE
        # -------------------------------------------------

        status_text.empty()


        # -------------------------------------------------
        # FINAL MESSAGE
        # -------------------------------------------------

        if st.session_state.processed_files:

            st.success(

                f"✅ "
                f"{len(st.session_state.processed_files)} "
                f"profile(s) processed successfully."

            )

        else:

            st.error(
                "No profiles could be processed."
            )


# =========================================================
# DISPLAY GENERATED FILES
#
# IMPORTANT:
# This is OUTSIDE the process button block.
#
# Therefore downloads remain visible even after
# downloading another file.
# =========================================================

if st.session_state.processed_files:


    st.divider()


    st.header(
        "📄 Updated Trainer Profiles"
    )


    # =====================================================
    # SHOW SKILLS APPLIED
    # =====================================================

    if st.session_state.skills_processed:

        st.write(
            "**Skills Applied:**"
        )


        st.info(

            "■ "
            + "    ■ ".join(
                st.session_state.skills_processed
            )

        )


    # =====================================================
    # UPDATE STATUS
    # =====================================================

    if st.session_state.status_rows:

        st.subheader(
            "Update Status"
        )


        st.dataframe(

            st.session_state.status_rows,

            use_container_width=True,

            hide_index=True

        )


    # =====================================================
    # INDIVIDUAL DOWNLOADS
    # =====================================================

    st.subheader(
        "Download Individual Profiles"
    )


    st.caption(
        "Downloading one profile will not remove "
        "the remaining download options."
    )


    for index, file_info in enumerate(
        st.session_state.processed_files
    ):


        col_file, col_download = (
            st.columns(
                [5, 2]
            )
        )


        with col_file:

            st.write(

                f"**{index + 1}. "
                f"{file_info['name']}**"

            )


        with col_download:

            st.download_button(

                label="⬇️ Download",

                data=file_info[
                    "data"
                ],

                file_name=file_info[
                    "name"
                ],

                mime=(
                    "application/vnd."
                    "openxmlformats-officedocument."
                    "wordprocessingml.document"
                ),

                key=(
                    f"download_profile_"
                    f"{index}"
                ),

                on_click="ignore",

                use_container_width=True

            )


    # =====================================================
    # CREATE ZIP FILE
    # =====================================================

    zip_buffer = BytesIO()


    with zipfile.ZipFile(

        zip_buffer,

        "w",

        zipfile.ZIP_DEFLATED

    ) as zip_file:


        for file_info in (
            st.session_state.processed_files
        ):


            zip_file.writestr(

                file_info["name"],

                file_info["data"]

            )


    zip_buffer.seek(0)


    # =====================================================
    # DOWNLOAD ALL AS ZIP
    # =====================================================

    st.divider()


    st.subheader(
        "Download All Profiles"
    )


    st.download_button(

        label=(
            "📦 Download All Updated "
            "Profiles as ZIP"
        ),

        data=zip_buffer.getvalue(),

        file_name=(
            "Updated_Trainer_Profiles.zip"
        ),

        mime="application/zip",

        type="primary",

        use_container_width=True,

        key="download_all_updated_profiles",

        on_click="ignore"

    )


    # =====================================================
    # CLEAR GENERATED RESULTS
    # =====================================================

    st.divider()


    if st.button(
        "🗑️ Clear Generated Files",
        use_container_width=True
    ):


        st.session_state.processed_files = []

        st.session_state.status_rows = []

        st.session_state.skills_processed = []


        st.rerun()